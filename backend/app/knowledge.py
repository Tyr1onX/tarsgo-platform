import base64
import hashlib
import io
import json
import logging
import os
import re
import stat
import uuid
import zipfile
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path, PurePosixPath
from typing import Protocol
from urllib.error import HTTPError, URLError
from urllib.parse import quote, urlsplit
from urllib.request import Request, urlopen

from docx import Document
from pypdf import PdfReader
from sqlalchemy import or_, select
from sqlalchemy.orm import Session, load_only

from .models import KnowledgeDocument


logger = logging.getLogger(__name__)
# PDF parser diagnostics may include snippets from malformed private documents.
_pypdf_logger = logging.getLogger("pypdf")
_pypdf_logger.addHandler(logging.NullHandler())
_pypdf_logger.propagate = False
ALLOWED_EXTENSIONS = {".md", ".txt", ".docx", ".pdf"}
MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_GITHUB_SYNC_FILES = 200
MAX_GITHUB_SYNC_BYTES = 50 * 1024 * 1024
MAX_EXTRACTED_TEXT_CHARS = 200_000
MAX_CURRENT_CONTEXT_CHARS = 5_000
MAX_HISTORY_CONTEXT_CHARS = 3_000
MAX_TOTAL_KNOWLEDGE_CONTEXT_CHARS = 8_000
MAX_HISTORY_RESULTS = 6
MAX_SEARCH_CANDIDATES = 60

_IGNORED_PATH_PARTS = {
    ".git",
    "node_modules",
    "vendor",
    ".venv",
    "venv",
    "__pycache__",
    "dist",
    "build",
    "coverage",
    "target",
    "backups",
    "database",
    "databases",
}
_SENSITIVE_NAME_PARTS = ("secret", "credential", "password", "private", "token", ".env")
_STOP_TERMS = {
    "一个", "一些", "今天", "需要", "准备", "活动", "事项", "团队", "负责", "安排", "相关", "可以",
    "进行", "然后", "目前", "本次", "历史", "资料", "文件", "经验", "以及", "我们", "他们",
}


class KnowledgeError(Exception):
    pass


class KnowledgeConfigError(KnowledgeError):
    pass


class KnowledgeSourceError(KnowledgeError):
    pass


class GitHubPathNotFound(KnowledgeSourceError):
    """A configured GitHub contents path does not exist at the requested ref."""


class UploadRejected(KnowledgeError):
    def __init__(self, status_code: int, detail: str):
        super().__init__(detail)
        self.status_code = status_code
        self.detail = detail


@dataclass(frozen=True)
class ExtractedText:
    content_text: str
    parse_status: str
    parse_error: str | None = None


@dataclass(frozen=True)
class KnowledgeContext:
    current_event_text: str = ""
    historical_text: str = ""
    current_event_documents: tuple["KnowledgeReference", ...] = ()
    historical_documents: tuple["KnowledgeReference", ...] = ()

    @property
    def context_chars(self) -> int:
        return len(self.current_event_text) + len(self.historical_text)


@dataclass(frozen=True)
class KnowledgeReference:
    id: int
    source_type: str
    source_name: str
    source_label: str
    title: str


def _knowledge_reference(document: KnowledgeDocument) -> KnowledgeReference:
    path = document.source_path or document.source_name
    title = (document.title or "").strip() or PurePosixPath(path).name or document.source_name
    if document.source_type == "github":
        path_parts = [part for part in path.split("/") if part]
        source_group = path_parts[0] if len(path_parts) > 1 else document.source_name.rsplit("/", 1)[-1]
        source_label = f"{source_group} · GitHub"
    else:
        source_label = f"{document.source_name} · 上传"
    return KnowledgeReference(
        id=document.id,
        source_type=document.source_type,
        source_name=document.source_name,
        source_label=source_label,
        title=title[:200],
    )


class GitHubReader(Protocol):
    def commit_sha(self, branch: str) -> str: ...
    def list_directory(self, path: str, ref: str) -> list[dict]: ...
    def read_blob(self, sha: str) -> bytes: ...


def knowledge_enabled() -> bool:
    return os.getenv("KNOWLEDGE_ENABLED", "false").strip().lower() in {"1", "true", "yes"}


def _utcnow() -> datetime:
    return datetime.utcnow()


def _source_key_hash(source_type: str, source_path: str) -> str:
    return hashlib.sha256(f"{source_type}\0{source_path}".encode("utf-8")).hexdigest()


def _repo_slug(repo: str) -> str:
    value = repo.strip()
    if value.startswith(("https://", "http://")):
        try:
            parsed = urlsplit(value)
            port = parsed.port
        except ValueError:
            raise KnowledgeConfigError("GitHub 仓库配置无效") from None
        if (
            parsed.scheme != "https"
            or parsed.hostname != "github.com"
            or parsed.username
            or parsed.password
            or port is not None
            or parsed.query
            or parsed.fragment
        ):
            raise KnowledgeConfigError("GitHub 仓库配置无效")
        value = parsed.path.strip("/")
    value = value.removesuffix(".git").strip("/")
    if not re.fullmatch(r"[A-Za-z0-9_.-]+/[A-Za-z0-9_.-]+", value):
        raise KnowledgeConfigError("GitHub 仓库配置无效")
    return value.lower()


def _github_paths() -> tuple[str, ...]:
    raw = os.getenv("KNOWLEDGE_GITHUB_PATHS", "docs,knowledge").strip() or "docs,knowledge"
    paths: list[str] = []
    for item in raw.split(","):
        candidate = item.strip().replace("\\", "/").strip("/")
        parts = PurePosixPath(candidate).parts
        if not candidate or candidate.startswith("/") or any(part in {".", ".."} for part in parts):
            raise KnowledgeConfigError("GitHub 允许目录配置无效")
        if candidate not in paths:
            paths.append(candidate)
    return tuple(paths)


def _path_is_ignored(path: str) -> bool:
    parts = PurePosixPath(path).parts
    if any(part.startswith(".") or part.lower() in _IGNORED_PATH_PARTS for part in parts):
        return True
    lowered = path.casefold()
    return any(value in lowered for value in _SENSITIVE_NAME_PARTS)


def _path_is_allowed(path: str, prefixes: tuple[str, ...]) -> bool:
    normalized = PurePosixPath(path).as_posix()
    if (
        normalized.startswith("/")
        or any(part in {".", ".."} for part in PurePosixPath(normalized).parts)
        or _path_is_ignored(normalized)
    ):
        return False
    suffix = PurePosixPath(normalized).suffix.lower()
    if suffix not in ALLOWED_EXTENSIONS:
        return False
    return any(normalized == prefix or normalized.startswith(prefix.rstrip("/") + "/") for prefix in prefixes)


def _decode_text_document(data: bytes) -> str:
    text = data.decode("utf-8-sig")
    if "\x00" in text:
        raise ValueError
    return text


def _extract_docx(data: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        entries = archive.infolist()
        if len(entries) > 2_000 or sum(item.file_size for item in entries) > 50 * 1024 * 1024:
            raise ValueError
        if "[Content_Types].xml" not in archive.namelist():
            raise ValueError
    document = Document(io.BytesIO(data))
    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _extract_pdf(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data), strict=False)
    if len(reader.pages) > 500:
        raise ValueError
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def extract_document_text(filename: str, data: bytes) -> ExtractedText:
    suffix = Path(filename).suffix.lower()
    try:
        if suffix in {".md", ".txt"}:
            text = _decode_text_document(data)
        elif suffix == ".docx":
            text = _extract_docx(data)
        elif suffix == ".pdf":
            text = _extract_pdf(data)
        else:
            raise UploadRejected(415, "仅支持 md、txt、docx、pdf 文件")
    except UploadRejected:
        raise
    except UnicodeDecodeError:
        return ExtractedText("", "failed", "文本编码无法识别")
    except Exception:
        return ExtractedText("", "failed", "文件无法解析")

    text = text.replace("\x00", "").strip()
    if not text:
        error = "未提取到可搜索文字（未执行 OCR）" if suffix == ".pdf" else "文件中没有可搜索文字"
        return ExtractedText("", "unparseable", error)
    if len(text) > MAX_EXTRACTED_TEXT_CHARS:
        return ExtractedText(
            text[:MAX_EXTRACTED_TEXT_CHARS],
            "truncated",
            f"提取文本超过 {MAX_EXTRACTED_TEXT_CHARS} 字，已截断",
        )
    return ExtractedText(text, "ready")


def _safe_upload_name(filename: str | None) -> str:
    value = (filename or "").replace("\\", "/").split("/")[-1]
    value = value.replace("\x00", "").replace("\r", "").replace("\n", "").strip()
    if not value or value in {".", ".."}:
        raise UploadRejected(400, "文件名无效")
    suffix = Path(value).suffix.lower()
    if suffix not in ALLOWED_EXTENSIONS:
        raise UploadRejected(415, "仅支持 md、txt、docx、pdf 文件")
    if len(value) > 255:
        stem = Path(value).stem[: 255 - len(suffix)]
        value = f"{stem}{suffix}"
    return value


def private_storage_dir(override: str | Path | None = None) -> Path:
    raw = override if override is not None else os.getenv("KNOWLEDGE_STORAGE_DIR", "/opt/tarsgo-knowledge")
    root = Path(raw).expanduser().resolve()
    module_path = Path(__file__).resolve()
    repository = module_path.parents[2] if module_path.parents[1].name == "backend" else module_path.parents[1]
    if root == repository or repository in root.parents:
        raise KnowledgeConfigError("知识文件存储目录必须位于仓库之外")
    root.mkdir(parents=True, exist_ok=True, mode=0o700)
    root.chmod(0o700)
    return root


def create_uploaded_document(
    db: Session,
    *,
    created_by: int,
    filename: str | None,
    data: bytes,
    storage_dir: str | Path | None = None,
) -> KnowledgeDocument:
    if len(data) > MAX_UPLOAD_BYTES:
        raise UploadRejected(413, "单个文件不能超过 10MB")
    safe_name = _safe_upload_name(filename)
    extraction = extract_document_text(safe_name, data)
    root = private_storage_dir(storage_dir)
    extension = Path(safe_name).suffix.lower()
    relative_path = f"uploads/{uuid.uuid4().hex}{extension}"
    upload_dir = root / "uploads"
    upload_dir.mkdir(mode=0o700, exist_ok=True)
    if upload_dir.is_symlink() or upload_dir.resolve().parent != root:
        raise KnowledgeConfigError("知识文件存储目录无效")
    upload_dir.chmod(0o700)
    target = upload_dir / Path(relative_path).name
    flags = os.O_WRONLY | os.O_CREAT | os.O_EXCL
    if hasattr(os, "O_NOFOLLOW"):
        flags |= os.O_NOFOLLOW
    descriptor = os.open(target, flags, stat.S_IRUSR | stat.S_IWUSR)
    try:
        with os.fdopen(descriptor, "wb") as stream:
            stream.write(data)
            stream.flush()
            os.fsync(stream.fileno())
        target.chmod(0o600)
        row = KnowledgeDocument(
            source_type="upload",
            source_name=safe_name,
            source_path=relative_path,
            source_key_hash=_source_key_hash("upload", relative_path),
            title=Path(safe_name).stem[:255] or safe_name,
            content_text=extraction.content_text,
            content_hash=hashlib.sha256(data).hexdigest(),
            git_commit_sha=None,
            source_updated_at=None,
            synced_at=_utcnow(),
            created_by=created_by,
            parse_status=extraction.parse_status,
            parse_error=extraction.parse_error,
            is_active=True,
        )
        db.add(row)
        db.commit()
        db.refresh(row)
        return row
    except BaseException:
        db.rollback()
        try:
            target.unlink()
        except FileNotFoundError:
            pass
        raise


def delete_document(db: Session, document: KnowledgeDocument, storage_dir: str | Path | None = None) -> None:
    relative_path = document.source_path if document.source_type == "upload" else None
    db.delete(document)
    db.commit()
    if relative_path:
        try:
            root = private_storage_dir(storage_dir)
            relative = PurePosixPath(relative_path)
            if relative.parts[:1] == ("uploads",) and len(relative.parts) == 2:
                target = root / relative_path
                if target.resolve().parent == (root / "uploads").resolve():
                    target.unlink(missing_ok=True)
        except Exception as exc:
            logger.warning(
                "knowledge upload file cleanup failed document_id=%s exception_type=%s",
                document.id,
                type(exc).__name__,
            )


class ReadOnlyGitHubClient:
    def __init__(self, repository: str, token: str | None = None):
        self.repository = _repo_slug(repository)
        self.token = token.strip() if token else ""
        if "\r" in self.token or "\n" in self.token:
            raise KnowledgeConfigError("GitHub 访问凭据配置无效")
        owner, name = self.repository.split("/", 1)
        self.api_root = f"https://api.github.com/repos/{quote(owner, safe='')}/{quote(name, safe='')}"

    def _json(self, endpoint: str, *, missing_path_is_ok: bool = False) -> dict | list:
        headers = {
            "Accept": "application/vnd.github+json",
            "User-Agent": "TARS-Go-Knowledge-Sync",
            "X-GitHub-Api-Version": "2022-11-28",
        }
        if self.token:
            headers["Authorization"] = f"Bearer {self.token}"
        request = Request(f"{self.api_root}{endpoint}", headers=headers)
        try:
            with urlopen(request, timeout=20) as response:
                payload = response.read(15 * 1024 * 1024 + 1)
            if len(payload) > 15 * 1024 * 1024:
                raise KnowledgeSourceError("GitHub 响应超过限制")
            decoded = json.loads(payload)
            if not isinstance(decoded, (dict, list)):
                raise KnowledgeSourceError("GitHub 响应无效")
            return decoded
        except HTTPError as exc:
            if exc.code == 404 and missing_path_is_ok:
                raise GitHubPathNotFound("GitHub 允许目录不存在") from None
            raise KnowledgeSourceError("GitHub 仓库暂时无法读取") from None
        except (URLError, TimeoutError, OSError, json.JSONDecodeError):
            raise KnowledgeSourceError("GitHub 仓库暂时无法读取") from None

    def commit_sha(self, branch: str) -> str:
        data = self._json(f"/commits/{quote(branch, safe='')}")
        sha = data.get("sha") if isinstance(data, dict) else None
        if not isinstance(sha, str) or not re.fullmatch(r"[0-9a-fA-F]{40,64}", sha):
            raise KnowledgeSourceError("GitHub 提交信息无效")
        return sha

    def list_directory(self, path: str, ref: str) -> list[dict]:
        encoded_path = quote(path, safe="/")
        data = self._json(
            f"/contents/{encoded_path}?ref={quote(ref, safe='')}&per_page=1000",
            missing_path_is_ok=True,
        )
        if isinstance(data, dict):
            return [data]
        if not all(isinstance(entry, dict) for entry in data):
            raise KnowledgeSourceError("GitHub 文件列表无效")
        return data

    def read_blob(self, sha: str) -> bytes:
        if not re.fullmatch(r"[0-9a-fA-F]{40,64}", sha):
            raise KnowledgeSourceError("GitHub 文件引用无效")
        data = self._json(f"/git/blobs/{sha}")
        if not isinstance(data, dict) or data.get("encoding") != "base64" or not isinstance(data.get("content"), str):
            raise KnowledgeSourceError("GitHub 文件内容无效")
        try:
            return base64.b64decode(data["content"], validate=False)
        except (ValueError, TypeError):
            raise KnowledgeSourceError("GitHub 文件内容无效") from None


def _allowed_github_entries(client: GitHubReader, commit_sha: str, prefixes: tuple[str, ...]) -> list[dict]:
    pending = list(prefixes)
    visited: set[str] = set()
    files: dict[str, dict] = {}
    while pending:
        path = pending.pop()
        if path in visited:
            continue
        visited.add(path)
        try:
            entries = client.list_directory(path, commit_sha)
        except GitHubPathNotFound:
            if path in prefixes:
                continue
            raise
        if len(entries) >= 1_000:
            raise KnowledgeSourceError("GitHub 单个目录条目超过同步上限")
        for entry in entries:
            entry_path = entry.get("path")
            entry_type = entry.get("type")
            if not isinstance(entry_path, str):
                continue
            if entry_type == "dir":
                under_allowed = any(
                    entry_path == prefix or entry_path.startswith(prefix.rstrip("/") + "/")
                    for prefix in prefixes
                )
                safe_parts = all(part not in {".", ".."} for part in PurePosixPath(entry_path).parts)
                if under_allowed and safe_parts and not _path_is_ignored(entry_path):
                    pending.append(entry_path)
            elif entry_type == "file" and _path_is_allowed(entry_path, prefixes):
                files[entry_path] = entry
                if len(files) > MAX_GITHUB_SYNC_FILES:
                    raise KnowledgeSourceError("允许目录内文件数量超过同步上限")
    return list(files.values())


def _upsert_github_document(
    db: Session,
    *,
    repository: str,
    path: str,
    commit_sha: str,
    blob_sha: str,
    content_hash: str,
    extraction: ExtractedText,
) -> bool:
    key_hash = _source_key_hash("github", f"{repository}:{path}")
    row = db.scalar(
        select(KnowledgeDocument)
        .options(load_only(
            KnowledgeDocument.id,
            KnowledgeDocument.source_type,
            KnowledgeDocument.source_name,
            KnowledgeDocument.source_path,
            KnowledgeDocument.source_key_hash,
            KnowledgeDocument.title,
            KnowledgeDocument.content_hash,
            KnowledgeDocument.source_blob_sha,
            KnowledgeDocument.git_commit_sha,
            KnowledgeDocument.source_updated_at,
            KnowledgeDocument.synced_at,
            KnowledgeDocument.created_by,
            KnowledgeDocument.parse_status,
            KnowledgeDocument.parse_error,
            KnowledgeDocument.is_active,
        ))
        .where(KnowledgeDocument.source_key_hash == key_hash)
    )
    is_new = row is None
    if row is None:
        row = KnowledgeDocument(
            source_type="github",
            source_name=repository,
            source_path=path,
            source_key_hash=key_hash,
            title=Path(path).stem[:255] or Path(path).name[:255],
            content_text="",
            content_hash=content_hash,
            git_commit_sha=commit_sha,
            source_updated_at=None,
            synced_at=_utcnow(),
            created_by=None,
            parse_status=extraction.parse_status,
            parse_error=extraction.parse_error,
            is_active=True,
        )
        db.add(row)
    row.source_name = repository
    row.source_path = path
    row.title = Path(path).stem[:255] or Path(path).name[:255]
    row.content_text = extraction.content_text
    row.content_hash = content_hash
    row.source_blob_sha = blob_sha
    row.git_commit_sha = commit_sha
    row.source_updated_at = None
    row.synced_at = _utcnow()
    row.created_by = None
    row.parse_status = extraction.parse_status
    row.parse_error = extraction.parse_error
    row.is_active = True
    return is_new


def sync_github_documents(
    db: Session,
    *,
    repository: str,
    branch: str,
    token: str | None = None,
    client: GitHubReader | None = None,
) -> dict[str, int]:
    repo = _repo_slug(repository)
    if not branch.strip() or len(branch) > 255:
        raise KnowledgeConfigError("GitHub 分支配置无效")
    prefixes = _github_paths()
    reader = client or ReadOnlyGitHubClient(repo, token)
    counts = {"added": 0, "updated": 0, "unchanged": 0, "failed": 0, "removed": 0}
    try:
        commit_sha = reader.commit_sha(branch)
        entries = _allowed_github_entries(reader, commit_sha, prefixes)
        total_size = sum(max(0, int(entry.get("size") or 0)) for entry in entries)
        if total_size > MAX_GITHUB_SYNC_BYTES:
            raise KnowledgeSourceError("GitHub 允许文件总大小超过同步上限")
        seen_paths = {entry["path"] for entry in entries}
        downloaded_bytes = 0
        for entry in entries:
            path = entry["path"]
            blob_sha = entry.get("sha")
            if not isinstance(blob_sha, str) or not re.fullmatch(r"[0-9a-fA-F]{40,64}", blob_sha):
                counts["failed"] += 1
                continue
            key_hash = _source_key_hash("github", f"{repo}:{path}")
            existing = db.scalar(
                select(KnowledgeDocument)
                .options(load_only(
                    KnowledgeDocument.id,
                    KnowledgeDocument.source_blob_sha,
                    KnowledgeDocument.git_commit_sha,
                    KnowledgeDocument.synced_at,
                    KnowledgeDocument.is_active,
                ))
                .where(KnowledgeDocument.source_key_hash == key_hash)
            )
            if existing and existing.is_active and existing.source_blob_sha == blob_sha:
                existing.git_commit_sha = commit_sha
                existing.synced_at = _utcnow()
                counts["unchanged"] += 1
                continue

            size = int(entry.get("size") or 0)
            if size < 0 or size > MAX_UPLOAD_BYTES:
                sentinel = hashlib.sha256(f"oversize:{blob_sha}:{size}".encode()).hexdigest()
                extraction = ExtractedText("", "failed", "文件超过 10MB 上限")
                is_new = _upsert_github_document(
                    db,
                    repository=repo,
                    path=path,
                    commit_sha=commit_sha,
                    blob_sha=blob_sha,
                    content_hash=sentinel,
                    extraction=extraction,
                )
                counts["added" if is_new else "updated"] += 1
                counts["failed"] += 1
                continue

            try:
                content = reader.read_blob(blob_sha)
            except KnowledgeSourceError:
                counts["failed"] += 1
                continue
            downloaded_bytes += len(content)
            if downloaded_bytes > MAX_GITHUB_SYNC_BYTES:
                raise KnowledgeSourceError("GitHub 文件实际下载总量超过同步上限")
            if len(content) > MAX_UPLOAD_BYTES:
                content = b""
                extraction = ExtractedText("", "failed", "文件超过 10MB 上限")
                content_hash = hashlib.sha256(f"oversize:{blob_sha}:{size}".encode()).hexdigest()
            else:
                content_hash = hashlib.sha256(content).hexdigest()
                extraction = extract_document_text(path, content)
            is_new = _upsert_github_document(
                db,
                repository=repo,
                path=path,
                commit_sha=commit_sha,
                blob_sha=blob_sha,
                content_hash=content_hash,
                extraction=extraction,
            )
            counts["added" if is_new else "updated"] += 1
            if extraction.parse_status in {"failed", "unparseable"}:
                counts["failed"] += 1

        existing_docs = db.scalars(
            select(KnowledgeDocument).options(load_only(
                KnowledgeDocument.id,
                KnowledgeDocument.source_path,
                KnowledgeDocument.is_active,
            )).where(
                KnowledgeDocument.source_type == "github",
                KnowledgeDocument.source_name == repo,
                KnowledgeDocument.is_active.is_(True),
            )
        ).all()
        for document in existing_docs:
            if (
                any(document.source_path == prefix or document.source_path.startswith(prefix.rstrip("/") + "/") for prefix in prefixes)
                and document.source_path not in seen_paths
            ):
                document.is_active = False
                document.parse_status = "removed"
                document.parse_error = None
                document.synced_at = _utcnow()
                counts["removed"] += 1
        db.commit()
    except BaseException:
        db.rollback()
        raise

    logger.info(
        "knowledge github sync completed added=%s updated=%s unchanged=%s failed=%s removed=%s",
        counts["added"], counts["updated"], counts["unchanged"], counts["failed"], counts["removed"],
    )
    return counts


def _search_terms(text: str) -> list[str]:
    lowered = text.casefold()
    terms: list[str] = []
    seen: set[str] = set()
    for word in re.findall(r"[a-z0-9][a-z0-9_-]{1,}", lowered):
        if word not in seen:
            terms.append(word)
            seen.add(word)
    for run in re.findall(r"[\u4e00-\u9fff]{2,}", lowered):
        for index in range(len(run) - 1):
            word = run[index : index + 2]
            if word not in _STOP_TERMS and word not in seen:
                terms.append(word)
                seen.add(word)
    return terms[:24]


def _planner_search_query(item_title: str | None, description: str, current_event_context: str) -> str:
    """Interleave terms so long input in one section cannot crowd the others out."""
    sections = [
        _search_terms(item_title or ""),
        _search_terms(description),
        _search_terms(current_event_context),
    ]
    terms: list[str] = []
    seen: set[str] = set()
    max_section_terms = max((len(section) for section in sections), default=0)
    for index in range(max_section_terms):
        for section in sections:
            if index >= len(section):
                continue
            term = section[index]
            if term in seen:
                continue
            terms.append(term)
            seen.add(term)
            if len(terms) >= 24:
                return " ".join(terms)
    return " ".join(terms)


def _relevance_score(document: KnowledgeDocument, terms: list[str]) -> int:
    title = (document.title or "").casefold()
    path = (document.source_path or "").casefold()
    source_name = (document.source_name or "").casefold()
    content = (document.content_text or "").casefold()
    score = 0
    for term in terms:
        if term in title:
            score += 8
        if term in path:
            score += 5
        if term in source_name:
            score += 3
        occurrences = content.count(term)
        score += min(occurrences, 3)
    return score


def _excerpt(content: str, terms: list[str], max_chars: int = 1_200) -> str:
    lowered = content.casefold()
    positions = [lowered.find(term) for term in terms]
    positions = [position for position in positions if position >= 0]
    start_at = min(positions, default=0)
    start = max(0, start_at - max_chars // 4)
    end = min(len(content), start + max_chars)
    start = max(0, end - max_chars)
    excerpt = content[start:end].strip()
    if start:
        excerpt = "…" + excerpt
    if end < len(content):
        excerpt += "…"
    return excerpt


def search_historical_documents(
    db: Session,
    query: str,
    *,
    excluded_ids: set[int] | None = None,
    max_chars: int = MAX_HISTORY_CONTEXT_CHARS,
) -> tuple[str, tuple[KnowledgeReference, ...]]:
    terms = _search_terms(query)
    if not terms or max_chars <= 0:
        return "", ()
    predicates = []
    for term in terms:
        pattern = f"%{term}%"
        predicates.extend(
            (
                KnowledgeDocument.title.ilike(pattern),
                KnowledgeDocument.source_path.ilike(pattern),
                KnowledgeDocument.content_text.ilike(pattern),
            )
        )
    statement = select(KnowledgeDocument).where(
        KnowledgeDocument.is_active.is_(True),
        KnowledgeDocument.parse_status.in_(("ready", "truncated")),
        or_(*predicates),
    )
    if excluded_ids:
        statement = statement.where(KnowledgeDocument.id.not_in(excluded_ids))
    candidates = db.scalars(statement.order_by(KnowledgeDocument.synced_at.desc()).limit(MAX_SEARCH_CANDIDATES)).all()
    ranked = [(score, document) for document in candidates if (score := _relevance_score(document, terms)) > 0]
    ranked.sort(key=lambda value: (value[0], value[1].synced_at), reverse=True)

    snippets: list[str] = []
    references: list[KnowledgeReference] = []
    used_chars = 0
    for _, document in ranked[:MAX_HISTORY_RESULTS]:
        reference = _knowledge_reference(document)
        header = f"资料《{reference.title[:120]}》\n"
        available = max_chars - used_chars - len(header)
        if available <= 0:
            break
        excerpt = _excerpt(document.content_text, terms, max_chars=min(1_200, available))
        if not excerpt:
            continue
        snippet = header + excerpt
        if len(snippet) > available:
            snippet = snippet[:available]
        snippets.append(snippet)
        references.append(reference)
        used_chars += len(snippet) + 2
    return "\n\n".join(snippets)[:max_chars], tuple(references)


def search_historical_knowledge(
    db: Session,
    query: str,
    *,
    excluded_ids: set[int] | None = None,
    max_chars: int = MAX_HISTORY_CONTEXT_CHARS,
) -> str:
    """Compatibility helper for callers that only need bounded history text."""
    text, _ = search_historical_documents(
        db, query, excluded_ids=excluded_ids, max_chars=max_chars
    )
    return text


def search_knowledge_documents(
    db: Session,
    query: str,
    *,
    limit: int = 20,
) -> list[KnowledgeReference]:
    """Search document metadata for the explicit current-event picker."""
    normalized = query.strip().casefold()
    if not normalized:
        return []
    terms = _search_terms(normalized) or [normalized]
    predicates = []
    for term in terms:
        pattern = f"%{term}%"
        predicates.extend(
            (
                KnowledgeDocument.title.ilike(pattern),
                KnowledgeDocument.source_path.ilike(pattern),
                KnowledgeDocument.source_name.ilike(pattern),
            )
        )
    rows = db.scalars(
        select(KnowledgeDocument)
        .where(
            KnowledgeDocument.is_active.is_(True),
            KnowledgeDocument.parse_status.in_(("ready", "truncated")),
            or_(*predicates),
        )
        .order_by(KnowledgeDocument.synced_at.desc(), KnowledgeDocument.id.desc())
        .limit(MAX_SEARCH_CANDIDATES)
    ).all()
    ranked: list[tuple[int, KnowledgeDocument]] = []
    for document in rows:
        title = (document.title or "").casefold()
        path = (document.source_path or "").casefold()
        source_name = (document.source_name or "").casefold()
        score = sum(
            (8 if term in title else 0)
            + (5 if term in path else 0)
            + (3 if term in source_name else 0)
            for term in terms
        )
        if normalized in title:
            score += 12
        elif normalized in path or normalized in source_name:
            score += 8
        if score:
            ranked.append((score, document))
    ranked.sort(key=lambda item: (item[0], item[1].synced_at), reverse=True)
    return [_knowledge_reference(document) for _, document in ranked[: max(1, min(limit, 20))]]


def build_planner_context(
    db: Session,
    *,
    description: str,
    item_title: str | None,
    current_event_context: str | None,
    current_event_document_ids: list[int],
    excluded_historical_document_ids: list[int] | None = None,
) -> KnowledgeContext:
    current_parts: list[str] = []
    pasted = (current_event_context or "").strip()
    if pasted:
        header = "负责人补充资料：\n"
        current_parts.append(header + pasted[: max(0, MAX_CURRENT_CONTEXT_CHARS - len(header))])
    used_current = sum(len(part) for part in current_parts) + max(0, len(current_parts) - 1) * 2

    selected_ids = list(dict.fromkeys(current_event_document_ids))[:5]
    selected_documents: list[KnowledgeDocument] = []
    if knowledge_enabled() and selected_ids:
        rows = db.scalars(
            select(KnowledgeDocument).where(
                KnowledgeDocument.id.in_(selected_ids),
                KnowledgeDocument.is_active.is_(True),
                KnowledgeDocument.parse_status.in_(("ready", "truncated")),
            )
        ).all()
        by_id = {row.id: row for row in rows}
        selected_documents = [by_id[document_id] for document_id in selected_ids if document_id in by_id]

    current_references: list[KnowledgeReference] = []
    for document in selected_documents:
        reference = _knowledge_reference(document)
        header = f"负责人选作本次资料《{reference.title[:100]}》：\n"
        available = MAX_CURRENT_CONTEXT_CHARS - used_current - len(header) - (2 if current_parts else 0)
        if available <= 0:
            break
        part = header + document.content_text[:available]
        current_parts.append(part)
        current_references.append(reference)
        used_current += len(part) + 2

    current_text = "\n\n".join(current_parts)[:MAX_CURRENT_CONTEXT_CHARS]
    remaining_history = min(MAX_HISTORY_CONTEXT_CHARS, MAX_TOTAL_KNOWLEDGE_CONTEXT_CHARS - len(current_text))
    history = ""
    historical_references: tuple[KnowledgeReference, ...] = ()
    if knowledge_enabled() and remaining_history > 0:
        query = _planner_search_query(item_title, description, pasted)
        history, historical_references = search_historical_documents(
            db,
            query,
            excluded_ids={document.id for document in selected_documents}
            | set(excluded_historical_document_ids or []),
            max_chars=remaining_history,
        )
    return KnowledgeContext(
        current_event_text=current_text,
        historical_text=history,
        current_event_documents=tuple(current_references),
        historical_documents=historical_references,
    )


def planner_input_text(description: str, item_title: str | None, context: KnowledgeContext) -> str:
    sections = [f"【负责人描述】\n{description.strip()}"]
    if item_title:
        sections.append(f"【事项标题提示】\n{item_title.strip()}")
    if context.current_event_text:
        sections.append(f"【本次事项资料】\n{context.current_event_text}")
    if context.historical_text:
        sections.append(f"【团队历史经验】\n{context.historical_text}")
    return "\n\n".join(sections)
