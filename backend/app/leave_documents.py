from __future__ import annotations

import io
import posixpath
import zipfile
from copy import deepcopy
from datetime import date
from pathlib import Path, PurePosixPath
from typing import Iterable

from docx import Document
from docx.shared import Inches
from lxml import etree

from .school_leave import (
    SCHOOL_LEAVE_DOCUMENT_REASON,
    format_school_leave_course_period,
    format_school_leave_time,
    school_leave_now,
)


DOCX_MIME_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
TEMPLATE_DIR = Path(__file__).with_name("templates")
DAILY_TEMPLATE_PATH = TEMPLATE_DIR / "school_leave.docx"
CAMP_TEMPLATE_PATH = TEMPLATE_DIR / "camp_leave_v2.docx"
SIGNATURE_PATH = TEMPLATE_DIR / "assets" / "instructor-signature.jpg"
SEAL_PATH = TEMPLATE_DIR / "assets" / "team-seal.jpg"

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
WP_NS = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
IMAGE_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"
XML_NS = "http://www.w3.org/XML/1998/namespace"
EMU_PER_POINT = 12_700


def _set_paragraph_text(paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def _find_paragraph(document, needle: str):
    return next((paragraph for paragraph in document.paragraphs if needle in paragraph.text), None)


def _append_image(paragraph, image_path: Path, width_inches: float) -> None:
    paragraph.add_run().add_picture(str(image_path), width=Inches(width_inches))


def _remove_image_parts(docx_bytes: bytes) -> bytes:
    """Remove electronic pictures while retaining blank signature/seal space."""
    source = zipfile.ZipFile(io.BytesIO(docx_bytes))
    files = {name: source.read(name) for name in source.namelist()}
    source.close()
    removed_targets: set[str] = set()
    parser = etree.XMLParser(resolve_entities=False, no_network=True)

    for name, data in list(files.items()):
        if name.endswith(".rels"):
            try:
                root = etree.fromstring(data, parser)
            except etree.XMLSyntaxError:
                continue
            rels_dir = PurePosixPath(name).parent
            source_dir = rels_dir.parent if rels_dir.name == "_rels" else rels_dir
            changed = False
            for rel in list(root):
                if rel.get("Type") != IMAGE_REL:
                    continue
                target = rel.get("Target", "")
                if target.startswith("/"):
                    part = posixpath.normpath(target.lstrip("/"))
                else:
                    part = posixpath.normpath(str(source_dir / target))
                removed_targets.add(str(PurePosixPath(part)))
                root.remove(rel)
                changed = True
            if changed:
                files[name] = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
        elif name.endswith(".xml"):
            try:
                root = etree.fromstring(data, parser)
            except etree.XMLSyntaxError:
                continue
            changed = False
            reserved_paragraph_heights: dict[etree._Element, int] = {}
            for node in list(root.iter(f"{{{W_NS}}}drawing")):
                parent = node.getparent()
                if parent is None:
                    continue
                paragraph = parent
                while paragraph is not None and paragraph.tag != f"{{{W_NS}}}p":
                    paragraph = paragraph.getparent()
                extent = node.find(f".//{{{WP_NS}}}extent")
                width = int(extent.get("cx", "0")) if extent is not None else 0
                height = int(extent.get("cy", "0")) if extent is not None else 0
                if paragraph is not None and width > 0:
                    # One 11pt blank space is roughly 5.5pt wide. The exact
                    # blank run keeps the original image's horizontal place.
                    spaces = max(1, round(width / (EMU_PER_POINT * 5.5)))
                    text = etree.Element(f"{{{W_NS}}}t")
                    text.set(f"{{{XML_NS}}}space", "preserve")
                    text.text = "\u00a0" * spaces
                    parent.replace(node, text)
                    reserved_paragraph_heights[paragraph] = max(
                        reserved_paragraph_heights.get(paragraph, 0), height
                    )
                else:
                    parent.remove(node)
                changed = True
            for node in list(root.iter(f"{{{W_NS}}}pict")):
                parent = node.getparent()
                if parent is not None:
                    parent.remove(node)
                    changed = True
            for paragraph, height in reserved_paragraph_heights.items():
                if paragraph.getroottree().getroot() is not root:
                    continue
                p_pr = paragraph.find(f"{{{W_NS}}}pPr")
                if p_pr is None:
                    p_pr = etree.Element(f"{{{W_NS}}}pPr")
                    paragraph.insert(0, p_pr)
                spacing = p_pr.find(f"{{{W_NS}}}spacing")
                if spacing is None:
                    spacing = etree.Element(f"{{{W_NS}}}spacing")
                    later_properties = {
                        "ind", "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc",
                        "textDirection", "textAlignment", "textboxTightWrap", "outlineLvl", "divId",
                        "cnfStyle", "rPr", "sectPr",
                    }
                    insertion_index = next(
                        (
                            index
                            for index, child in enumerate(p_pr)
                            if etree.QName(child).localname in later_properties
                        ),
                        len(p_pr),
                    )
                    p_pr.insert(insertion_index, spacing)
                spacing.set(f"{{{W_NS}}}line", str(max(1, round(height / EMU_PER_POINT * 20))))
                spacing.set(f"{{{W_NS}}}lineRule", "exact")
            if changed:
                files[name] = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)

    for target in removed_targets:
        files.pop(target, None)
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as result:
        for name, data in files.items():
            result.writestr(name, data)
    return output.getvalue()


def _save(document, *, offline: bool) -> bytes:
    output = io.BytesIO()
    document.save(output)
    content = output.getvalue()
    return _remove_image_parts(content) if offline else content


def build_daily_leave_v2_docx(
    *,
    start_at,
    end_at,
    name: str,
    student_id: str,
    college_name: str,
    contact_phone: str,
    offline: bool,
) -> bytes:
    """Render one personal daily leave proof from the current approved template."""
    if not contact_phone.strip():
        raise ValueError("LEAVE_CONTACT_PHONE 未配置，无法生成请假材料")
    document = Document(DAILY_TEMPLATE_PATH)
    body = _find_paragraph(document, "以下学生因")
    phone = _find_paragraph(document, "联系电话：")
    issued = next(
        (paragraph for paragraph in document.paragraphs if "2026年9月20日" in paragraph.text.replace(" ", "")),
        None,
    )
    if body is None or phone is None or issued is None or not document.tables:
        raise ValueError("日常请假模板结构不完整")

    body_text = (
        f"以下学生因参加{format_school_leave_time(start_at, end_at)}的"
        f"{SCHOOL_LEAVE_DOCUMENT_REASON}，不能参加"
        f"{format_school_leave_course_period(start_at, end_at)}，特此证明。"
    )
    _set_paragraph_text(body, body_text)
    phone_indent = phone.text[: len(phone.text) - len(phone.text.lstrip())]
    _set_paragraph_text(phone, f"{phone_indent}联系电话：{contact_phone.strip()}")
    issued_indent = issued.text[: len(issued.text) - len(issued.text.lstrip())]
    generated_at = school_leave_now()
    _set_paragraph_text(issued, f"{issued_indent}{generated_at.year}年{generated_at.month}月{generated_at.day}日")

    table = document.tables[0]
    if len(table.columns) == 2:
        table.add_column(Inches(1.4))
    if len(table.columns) != 3 or len(table.rows) < 2:
        raise ValueError("日常请假模板名单表格结构不完整")
    _set_paragraph_text(table.rows[0].cells[0].paragraphs[0], "姓名")
    _set_paragraph_text(table.rows[0].cells[1].paragraphs[0], "学号")
    _set_paragraph_text(table.rows[0].cells[2].paragraphs[0], "学院")
    _set_paragraph_text(table.rows[1].cells[0].paragraphs[0], name)
    _set_paragraph_text(table.rows[1].cells[1].paragraphs[0], student_id)
    _set_paragraph_text(table.rows[1].cells[2].paragraphs[0], college_name)

    if not offline:
        signature = _find_paragraph(document, "指导教师（签字）：")
        if signature is None:
            signature = _find_paragraph(document, "指导教师")
        if signature is None:
            raise ValueError("日常请假模板缺少指导教师签字位置")
        _append_image(signature, SIGNATURE_PATH, 0.89)
        _append_image(signature, SEAL_PATH, 1.10)
    return _save(document, offline=offline)


def _camp_date(value: date) -> str:
    return f"{value.year}年{value.month}月{value.day}日"


def build_camp_leave_college_docx(
    *,
    title: str,
    event_type: str,
    start_date: date,
    end_date: date,
    participants: Iterable[tuple[str, str, str]],
    offline: bool,
) -> bytes:
    """Render a single college's camp leave request from the supplied template."""
    if event_type not in {"winter", "summer"}:
        raise ValueError("集中请假活动类型无效")
    document = Document(CAMP_TEMPLATE_PATH)
    paragraphs = document.paragraphs
    body = _find_paragraph(document, "兹有以下学生")
    if body is None or not document.tables:
        raise ValueError("集中请假模板结构不完整")
    season = "寒假" if event_type == "winter" else "暑假"
    body_text = (
        f"兹有以下学生，因参与即将举行的“{title}”，需{season}留校进行机器人的设计、编程、测试，"
        f"时间是{_camp_date(start_date)}到{_camp_date(end_date)}。为确保留校期间安全、高效有序，"
        "吉甲大师双创基地已安排指导教师全程负责学生的日常管理，包括早晚签到，"
        "并进行全天候的实验室活动安排。"
    )
    _set_paragraph_text(body, body_text)

    table = document.tables[0]
    if len(table.rows) < 2 or len(table.columns) != 4:
        raise ValueError("集中请假模板名单表格结构不完整")
    row_template = deepcopy(table.rows[1]._tr)
    for row in list(table.rows)[1:]:
        table._tbl.remove(row._tr)
    rows = list(participants)
    if not rows:
        raise ValueError("该学院没有可生成的名单")
    for index, (name, student_id, college_name) in enumerate(rows, start=1):
        table._tbl.append(deepcopy(row_template))
        row = table.rows[-1]
        values = (str(index), name, student_id, college_name)
        for cell, value in zip(row.cells, values):
            _set_paragraph_text(cell.paragraphs[0], value)

    return _save(document, offline=offline)
