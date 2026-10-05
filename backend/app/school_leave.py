from __future__ import annotations

import io
import os
import re
import secrets
from copy import deepcopy
from dataclasses import dataclass
from datetime import datetime, timedelta
from pathlib import Path
from zoneinfo import ZoneInfo

from docx import Document
from sqlalchemy import select
from sqlalchemy.orm import Session

from .models import SchoolLeaveGroupResult, SchoolLeaveRequest, SchoolLeaveRun


DEFAULT_SCHOOL_LEAVE_REASON = "参加吉林大学吉甲大师机器人战队相关创新实践活动及工作安排"
SHANGHAI_TZ = ZoneInfo("Asia/Shanghai")
SCHOOL_LEAVE_TEMPLATE_PATH = Path(__file__).with_name("templates") / "school_leave.docx"
SCHOOL_LEAVE_RESULT_MAX_BYTES = 15 * 1024 * 1024
SCHOOL_LEAVE_RESULT_TTL = timedelta(hours=72)
SCHOOL_LEAVE_RESULT_MIME_BY_SUFFIX = {
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".png": "image/png",
    ".pdf": "application/pdf",
}
_CUTOFF_RE = re.compile(r"^(?:[01]\d|2[0-3]):[0-5]\d$")


@dataclass(frozen=True)
class SchoolLeaveGroup:
    index: int
    start_at: datetime
    end_at: datetime
    requests: tuple[SchoolLeaveRequest, ...]


def school_leave_now() -> datetime:
    return datetime.now(SHANGHAI_TZ).replace(tzinfo=None, second=0, microsecond=0)


def get_leave_contact_phone() -> str:
    return os.getenv("LEAVE_CONTACT_PHONE", "").strip()


def get_leave_daily_cutoff() -> str:
    value = os.getenv("LEAVE_DAILY_CUTOFF", "").strip()
    if not _CUTOFF_RE.fullmatch(value):
        raise ValueError("LEAVE_DAILY_CUTOFF 必须配置为 HH:MM 24 小时格式")
    return value


def get_school_leave_result_storage_dir() -> Path:
    return Path(
        os.getenv(
            "SCHOOL_LEAVE_RESULT_STORAGE_DIR",
            "/opt/tarsgo-data/school-leave-results",
        )
    )


def validate_school_leave_result_upload(
    filename: str | None,
    content_type: str | None,
    content: bytes,
) -> tuple[str, str]:
    original_filename = (filename or "").replace("\\", "/").split("/")[-1].strip()
    suffix = Path(original_filename).suffix.lower()
    expected_mime = SCHOOL_LEAVE_RESULT_MIME_BY_SUFFIX.get(suffix)
    if not original_filename or expected_mime is None or content_type != expected_mime:
        raise ValueError("仅支持 JPG、JPEG、PNG、PDF，且文件类型必须与扩展名一致")
    if len(content) > SCHOOL_LEAVE_RESULT_MAX_BYTES:
        raise OverflowError("盖章材料不能超过 15 MiB")
    return original_filename[:255], suffix


def store_school_leave_result_file(
    content: bytes,
    suffix: str,
    *,
    storage_dir: Path | None = None,
) -> str:
    directory = storage_dir or get_school_leave_result_storage_dir()
    directory.mkdir(parents=True, exist_ok=True)
    stored_name = f"{secrets.token_hex(24)}{suffix}"
    (directory / stored_name).write_bytes(content)
    return stored_name


def school_leave_result_file_path(
    stored_name: str,
    *,
    storage_dir: Path | None = None,
) -> Path:
    directory = storage_dir or get_school_leave_result_storage_dir()
    safe_name = Path(stored_name).name
    if safe_name != stored_name:
        raise ValueError("invalid stored school leave result name")
    return directory / safe_name


def delete_school_leave_result_file(
    stored_name: str,
    *,
    storage_dir: Path | None = None,
) -> None:
    school_leave_result_file_path(stored_name, storage_dir=storage_dir).unlink(missing_ok=True)


def school_leave_result_download_filename(run: SchoolLeaveRun, mime_type: str) -> str:
    extension = {
        "image/jpeg": ".jpg",
        "image/png": ".png",
        "application/pdf": ".pdf",
    }[mime_type]
    return f"吉甲大师请假条_{run.collected_at:%Y-%m-%d}_盖章{extension}"


def cleanup_expired_school_leave_results(
    db: Session,
    *,
    now: datetime | None = None,
    storage_dir: Path | None = None,
) -> int:
    cleanup_at = now or school_leave_now()
    results = list(
        db.scalars(
            select(SchoolLeaveGroupResult)
            .where(
                SchoolLeaveGroupResult.expires_at <= cleanup_at,
                SchoolLeaveGroupResult.deleted_at.is_(None),
            )
            .order_by(SchoolLeaveGroupResult.id)
        )
    )
    cleaned = 0
    for result in results:
        delete_school_leave_result_file(result.stored_name, storage_dir=storage_dir)
        result.deleted_at = cleanup_at
        db.commit()
        cleaned += 1
    return cleaned


def collect_pending_school_leave(
    db: Session,
    *,
    created_by: int | None,
    collected_at: datetime | None = None,
) -> SchoolLeaveRun | None:
    """Freeze every request that is pending at the collection transaction boundary."""
    try:
        pending = list(
            db.scalars(
                select(SchoolLeaveRequest)
                .where(SchoolLeaveRequest.status == "pending")
                .order_by(SchoolLeaveRequest.id)
                .with_for_update()
            )
        )
        if not pending:
            db.commit()
            return None

        run = SchoolLeaveRun(
            collected_at=collected_at or school_leave_now(),
            created_by=created_by,
            reason=DEFAULT_SCHOOL_LEAVE_REASON,
            status="ready",
        )
        db.add(run)
        db.flush()
        for request in pending:
            request.status = "included"
            request.run_id = run.id
        db.commit()
        db.refresh(run)
        return run
    except Exception:
        db.rollback()
        raise


def requests_for_run(db: Session, run_id: int) -> list[SchoolLeaveRequest]:
    return list(
        db.scalars(
            select(SchoolLeaveRequest)
            .where(
                SchoolLeaveRequest.run_id == run_id,
                SchoolLeaveRequest.status == "included",
            )
            .order_by(
                SchoolLeaveRequest.start_at,
                SchoolLeaveRequest.end_at,
                SchoolLeaveRequest.student_id_snapshot,
                SchoolLeaveRequest.member_name_snapshot,
                SchoolLeaveRequest.id,
            )
        )
    )


def group_school_leave_requests(requests: list[SchoolLeaveRequest]) -> list[SchoolLeaveGroup]:
    grouped: dict[tuple[datetime, datetime], list[SchoolLeaveRequest]] = {}
    for request in sorted(
        requests,
        key=lambda item: (
            item.start_at,
            item.end_at,
            item.student_id_snapshot,
            item.member_name_snapshot,
            item.id,
        ),
    ):
        grouped.setdefault((request.start_at, request.end_at), []).append(request)

    return [
        SchoolLeaveGroup(
            index=index,
            start_at=start_at,
            end_at=end_at,
            requests=tuple(items),
        )
        for index, ((start_at, end_at), items) in enumerate(grouped.items())
    ]


def format_school_leave_time(start_at: datetime, end_at: datetime) -> str:
    if start_at.date() == end_at.date():
        return (
            f"{start_at.year} 年 {start_at.month} 月 {start_at.day} 日 "
            f"{start_at:%H:%M} 至 {end_at:%H:%M}"
        )
    return (
        f"{start_at.year} 年 {start_at.month} 月 {start_at.day} 日 {start_at:%H:%M} 至\n"
        f"{end_at.year} 年 {end_at.month} 月 {end_at.day} 日 {end_at:%H:%M}"
    )


def _replace_paragraph_text(paragraph, text: str) -> None:
    if not paragraph.runs:
        paragraph.add_run(text)
        return
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def _fill_school_leave_template(
    document,
    run: SchoolLeaveRun,
    group: SchoolLeaveGroup,
    *,
    contact_phone: str,
) -> None:
    body = next(paragraph for paragraph in document.paragraphs if paragraph.text.startswith("以下学生因"))
    phone = next(paragraph for paragraph in document.paragraphs if "联系电话：" in paragraph.text)
    issued = next(
        paragraph
        for paragraph in document.paragraphs
        if "2026年9月20日" in paragraph.text.replace(" ", "")
    )
    table = document.tables[0]

    phone_indent = phone.text[: len(phone.text) - len(phone.text.lstrip())]
    issued_indent = issued.text[: len(issued.text) - len(issued.text.lstrip())]

    body_text = body.text.replace(
        "9月20日下午3点到5点",
        format_school_leave_time(group.start_at, group.end_at),
    ).replace(
        "吉甲大师双创基地参观接待活动",
        run.reason,
    )
    _replace_paragraph_text(body, body_text)
    _replace_paragraph_text(phone, f"{phone_indent}联系电话：{contact_phone.strip()}")
    _replace_paragraph_text(
        issued,
        f"{issued_indent}{run.collected_at.year}年{run.collected_at.month}月{run.collected_at.day}日",
    )

    row_template = deepcopy(table.rows[1]._tr)
    for index, request in enumerate(group.requests):
        if index == 0:
            row = table.rows[1]
        else:
            table._tbl.append(deepcopy(row_template))
            row = table.rows[-1]
        _replace_paragraph_text(row.cells[0].paragraphs[0], request.member_name_snapshot)
        _replace_paragraph_text(row.cells[1].paragraphs[0], request.student_id_snapshot)


def _school_leave_group_document(
    run: SchoolLeaveRun,
    group: SchoolLeaveGroup,
    *,
    contact_phone: str,
):
    document = Document(SCHOOL_LEAVE_TEMPLATE_PATH)
    _fill_school_leave_template(document, run, group, contact_phone=contact_phone)
    return document


def _append_school_leave_page(document, page) -> None:
    document.add_page_break()
    section_properties = document.element.body.sectPr
    for child in page.element.body:
        if child.tag.endswith("}sectPr"):
            continue
        section_properties.addprevious(deepcopy(child))


def build_school_leave_docx(
    run: SchoolLeaveRun,
    group: SchoolLeaveGroup,
    *,
    contact_phone: str,
) -> bytes:
    """Build one exact-time group DOCX from the approved School Leave template."""
    if not contact_phone.strip():
        raise ValueError("LEAVE_CONTACT_PHONE 未配置，无法生成学校请假材料")

    document = _school_leave_group_document(run, group, contact_phone=contact_phone)
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def build_school_leave_run_docx(
    run: SchoolLeaveRun,
    groups: list[SchoolLeaveGroup],
    *,
    contact_phone: str,
) -> bytes:
    """Build one template-backed DOCX, with each exact-time group on its own page."""
    if not contact_phone.strip():
        raise ValueError("LEAVE_CONTACT_PHONE 未配置，无法生成学校请假材料")
    if not groups:
        raise ValueError("该汇总批次没有可生成的请假材料")

    document = _school_leave_group_document(run, groups[0], contact_phone=contact_phone)
    for group in groups[1:]:
        page = _school_leave_group_document(run, group, contact_phone=contact_phone)
        _append_school_leave_page(document, page)

    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def school_leave_document_filename(group: SchoolLeaveGroup) -> str:
    start = group.start_at.strftime("%Y-%m-%d_%H%M")
    end = group.end_at.strftime("%Y-%m-%d_%H%M")
    return f"请假条_{start}-{end}.docx"


def school_leave_run_document_filename(db: Session, run: SchoolLeaveRun) -> str:
    date_text = run.collected_at.strftime("%Y-%m-%d")
    day_start = run.collected_at.replace(hour=0, minute=0, second=0, microsecond=0)
    day_end = day_start + timedelta(days=1)
    valid_run_ids = list(
        db.scalars(
            select(SchoolLeaveRun.id)
            .where(
                SchoolLeaveRun.status != "cancelled",
                SchoolLeaveRun.collected_at >= day_start,
                SchoolLeaveRun.collected_at < day_end,
            )
            .order_by(SchoolLeaveRun.collected_at, SchoolLeaveRun.id)
        )
    )
    supplement_index = valid_run_ids.index(run.id)
    supplement_suffix = "" if supplement_index == 0 else f"_补充{supplement_index}"
    return f"吉甲大师请假条_{date_text}{supplement_suffix}.docx"
