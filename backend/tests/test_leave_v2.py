"""MySQL-backed coverage for daily leave V2 and template-backed documents."""

from __future__ import annotations

import hashlib
import io
import os
import zipfile
from datetime import date, datetime, timedelta
from pathlib import Path
from uuid import uuid4
from zoneinfo import ZoneInfo

from docx import Document
from fastapi.testclient import TestClient
from sqlalchemy import delete, func, select

from app.auth import get_current_member
from app.db import SessionLocal, get_db
from app.leave_documents import build_camp_leave_college_docx
from app.main import app
from app.models import (
    CampLeaveEvent,
    CampLeaveParticipant,
    DailyLeaveEntry,
    DailyLeaveWindow,
    Member,
    SchoolLeaveRequest,
    SchoolLeaveRun,
)


TZ = ZoneInfo("Asia/Shanghai")


def local_now() -> datetime:
    return datetime.now(TZ).replace(tzinfo=None, second=0, microsecond=0)


def docx_text(content: bytes) -> tuple[list[str], list[list[str]]]:
    document = Document(io.BytesIO(content))
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    tables = [[[cell.text for cell in row.cells] for row in table.rows] for table in document.tables]
    return paragraphs, tables


def package_media(content: bytes) -> set[bytes]:
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        return {archive.read(name) for name in archive.namelist() if "/media/" in name or name.startswith("media/")}


def main() -> None:
    suffix = uuid4().hex
    admin_email = f"leave-v2-admin-{suffix}@example.invalid"
    formal_email = f"leave-v2-formal-{suffix}@example.invalid"
    reserve_email = f"leave-v2-reserve-{suffix}@example.invalid"
    unassigned_email = f"leave-v2-unassigned-{suffix}@example.invalid"
    incomplete_email = f"leave-v2-incomplete-{suffix}@example.invalid"
    event_ids: set[int] = set()
    window_ids: set[int] = set()
    legacy_run_ids: set[int] = set()
    previous_overrides = dict(app.dependency_overrides)
    original_phone = os.environ.get("LEAVE_CONTACT_PHONE")
    os.environ["LEAVE_CONTACT_PHONE"] = "000-0000-0000"

    with SessionLocal() as db:
        admin = Member(name="V2 测试管理员", email=admin_email, role="admin", status="active")
        formal = Member(
            name="正式成员甲",
            email=formal_email,
            student_id="26000201",
            college="electronic-science-engineering",
            team_membership="formal",
            role="member",
            status="active",
        )
        reserve = Member(
            name="Reserve Member",
            email=reserve_email,
            student_id="26000202",
            college="artificial-intelligence",
            team_membership="reserve",
            role="member",
            status="active",
        )
        unassigned = Member(
            name="身份待管理员维护",
            email=unassigned_email,
            student_id="26000204",
            college="art",
            team_membership=None,
            role="member",
            status="active",
        )
        incomplete = Member(
            name="资料未全",
            email=incomplete_email,
            student_id=None,
            college=None,
            team_membership="formal",
            role="member",
            status="active",
        )
        db.add_all([admin, formal, reserve, unassigned, incomplete])
        db.commit()
        db.refresh(admin)
        db.refresh(formal)
        db.refresh(reserve)
        db.refresh(unassigned)
        db.refresh(incomplete)

        def override_database():
            yield db

        app.dependency_overrides[get_db] = override_database
        actor: Member | None = None

        def set_actor(member: Member | None) -> None:
            nonlocal actor
            actor = member
            if member is None:
                app.dependency_overrides.pop(get_current_member, None)
            else:
                app.dependency_overrides[get_current_member] = lambda: actor

        client = TestClient(app)
        try:
            now = local_now()
            activity_start = (now + timedelta(days=3)).replace(hour=13, minute=30)
            activity_end = activity_start.replace(hour=17, minute=10)

            # Team members can generate directly without an administrator-created window.
            standalone_payload = {
                "start_at": activity_start.isoformat(timespec="minutes"),
                "end_at": activity_end.isoformat(timespec="minutes"),
            }
            set_actor(formal)
            standalone_doc = client.post("/api/daily-leave/self-service/document", json=standalone_payload)
            assert standalone_doc.status_code == 200, standalone_doc.text
            standalone_paragraphs, standalone_tables = docx_text(standalone_doc.content)
            assert any(
                "吉甲大师双创基地机器人战队创新实践活动" in line.replace("\u2060", "")
                for line in standalone_paragraphs
            )
            assert not any("内部活动" in line for line in standalone_paragraphs)
            assert any("下午课程" in line for line in standalone_paragraphs)
            assert standalone_tables[0][1] == ["正式成员甲", "26000201", "电子科学与工程学院"]
            standalone_entry = db.scalar(
                select(DailyLeaveEntry).where(
                    DailyLeaveEntry.window_id.is_(None), DailyLeaveEntry.member_id == formal.id
                ).order_by(DailyLeaveEntry.id)
            )
            assert standalone_entry is not None
            assert standalone_entry.start_at == activity_start.replace(second=0, microsecond=0)
            assert standalone_entry.end_at == activity_end.replace(second=0, microsecond=0)
            assert standalone_entry.name_snapshot == "正式成员甲"
            assert standalone_entry.student_id_snapshot == "26000201"
            assert standalone_entry.college_snapshot == "electronic-science-engineering"
            assert standalone_entry.participant_type == "formal"
            assert standalone_entry.submitted_at is not None
            standalone_offline = client.post(
                "/api/daily-leave/self-service/document?offline=true", json=standalone_payload
            )
            assert standalone_offline.status_code == 200
            assert package_media(standalone_offline.content) == set()
            assert client.post(
                "/api/daily-leave/self-service/document",
                json={**standalone_payload, "title": "不允许客户端提交活动名"},
            ).status_code == 422
            assert client.post(
                "/api/daily-leave/self-service/document",
                json={**standalone_payload, "end_at": standalone_payload["start_at"]},
            ).status_code == 422
            set_actor(unassigned)
            unassigned_self_doc = client.post("/api/daily-leave/self-service/document", json=standalone_payload)
            assert unassigned_self_doc.status_code == 200, unassigned_self_doc.text
            unassigned_self_entry = db.scalar(
                select(DailyLeaveEntry).where(
                    DailyLeaveEntry.window_id.is_(None), DailyLeaveEntry.member_id == unassigned.id
                )
            )
            assert unassigned_self_entry is not None and unassigned_self_entry.participant_type == "other"
            assert db.scalar(
                select(func.count(DailyLeaveEntry.id)).where(
                    DailyLeaveEntry.window_id.is_(None), DailyLeaveEntry.member_id == formal.id
                )
            ) == 2
            set_actor(incomplete)
            assert client.post("/api/daily-leave/self-service/document", json=standalone_payload).status_code == 409

            window_payload = {
                "title": "管理员内部留档名称",
                "start_at": activity_start.isoformat(timespec="minutes"),
                "end_at": activity_end.isoformat(timespec="minutes"),
            }

            set_actor(formal)
            assert client.post("/api/daily-leave/admin/windows", json=window_payload).status_code == 403
            set_actor(admin)
            invalid_window = {**window_payload, "end_at": window_payload["start_at"]}
            assert client.post("/api/daily-leave/admin/windows", json=invalid_window).status_code == 422
            created = client.post("/api/daily-leave/admin/windows", json=window_payload)
            assert created.status_code == 201, created.text
            window = created.json()
            window_id = window["id"]
            window_ids.add(window_id)
            assert window["open_until"] == window["end_at"]
            assert window["public_enabled"] is False
            assert window["public_path"] is None
            assert window["team_open"] is True
            set_actor(formal)
            assert client.post(f"/api/daily-leave/admin/windows/{window_id}/public-link").status_code == 403
            set_actor(admin)
            enabled_link = client.post(f"/api/daily-leave/admin/windows/{window_id}/public-link")
            assert enabled_link.status_code == 200, enabled_link.text
            window = enabled_link.json()
            token = window["public_path"].rsplit("/", 1)[1]
            assert window["public_enabled"] is True
            assert 40 <= len(token) <= 50
            assert set(token) <= set("abcdefghijklmnopqrstuvwxyzABCDEFGHIJKLMNOPQRSTUVWXYZ0123456789_-")

            # Formal and reserve members generate immediately from snapshots; an incomplete profile is rejected.
            set_actor(formal)
            member_windows = client.get("/api/daily-leave/windows")
            assert member_windows.status_code == 200
            assert member_windows.json()[0]["title"] == window_payload["title"]
            formal_doc = client.post(f"/api/daily-leave/windows/{window_id}/document")
            assert formal_doc.status_code == 200, formal_doc.text
            first_entry = db.scalar(select(DailyLeaveEntry).where(DailyLeaveEntry.window_id == window_id))
            assert first_entry is not None
            assert first_entry.participant_type == "formal"
            assert first_entry.name_snapshot == "正式成员甲"
            assert first_entry.student_id_snapshot == "26000201"
            assert first_entry.college_snapshot == "electronic-science-engineering"
            assert first_entry.start_at is None and first_entry.end_at is None
            assert client.post(f"/api/daily-leave/windows/{window_id}/document").status_code == 200
            assert db.scalar(select(func.count(DailyLeaveEntry.id)).where(DailyLeaveEntry.window_id == window_id)) == 1

            set_actor(reserve)
            reserve_doc = client.post(f"/api/daily-leave/windows/{window_id}/document")
            assert reserve_doc.status_code == 200, reserve_doc.text
            reserve_entry = db.scalar(
                select(DailyLeaveEntry).where(
                    DailyLeaveEntry.window_id == window_id,
                    DailyLeaveEntry.member_id == reserve.id,
                )
            )
            assert reserve_entry is not None and reserve_entry.participant_type == "reserve"

            set_actor(unassigned)
            unassigned_windows = client.get("/api/daily-leave/windows")
            assert unassigned_windows.status_code == 200
            assert any(item["id"] == window_id for item in unassigned_windows.json())
            unassigned_doc = client.post(f"/api/daily-leave/windows/{window_id}/document")
            assert unassigned_doc.status_code == 200, unassigned_doc.text
            unassigned_entry = db.scalar(
                select(DailyLeaveEntry).where(
                    DailyLeaveEntry.window_id == window_id,
                    DailyLeaveEntry.member_id == unassigned.id,
                )
            )
            assert unassigned_entry is not None and unassigned_entry.participant_type == "other"

            set_actor(incomplete)
            assert client.post(f"/api/daily-leave/windows/{window_id}/document").status_code == 409

            # Default daily DOCX embeds exact source media and deterministic body/date/identity; offline omits both images.
            signed_bytes = formal_doc.content
            signed_paragraphs, signed_tables = docx_text(signed_bytes)
            assert not any(window_payload["title"] in line for line in signed_paragraphs)
            time_text = (
                f"{activity_start.year}年{activity_start.month}月{activity_start.day}日"
                "13:30 至 17:10"
            )
            assert any(
                "吉甲大师双创基地机器人战队创新实践活动" in line.replace("\u2060", "")
                for line in signed_paragraphs
            )
            assert any("下午课程" in line for line in signed_paragraphs)
            assert any(time_text == line.removeprefix("以下学生因参加").split("的")[0] for line in signed_paragraphs if line.startswith("以下学生因参加"))
            assert signed_tables[0][0] == ["姓名", "学号", "学院"]
            assert signed_tables[0][1] == ["正式成员甲", "26000201", "电子科学与工程学院"]
            signature = (Path(__file__).parents[1] / "app/templates/assets/instructor-signature.jpg").read_bytes()
            seal = (Path(__file__).parents[1] / "app/templates/assets/team-seal.jpg").read_bytes()
            assert {hashlib.sha256(value).hexdigest() for value in package_media(signed_bytes)} >= {
                hashlib.sha256(signature).hexdigest(),
                hashlib.sha256(seal).hexdigest(),
            }
            set_actor(formal)
            offline = client.post(f"/api/daily-leave/windows/{window_id}/document?offline=true")
            assert offline.status_code == 200, offline.text
            assert package_media(offline.content) == set()
            offline_paragraphs, offline_tables = docx_text(offline.content)
            assert [paragraph.replace("\u00a0", "") for paragraph in offline_paragraphs] == signed_paragraphs
            assert offline_tables == signed_tables

            # Public link exposes only fixed activity fields and repeats the original participant snapshot.
            set_actor(None)
            public_info = client.get(f"/api/daily-leave/public/{token}")
            assert public_info.status_code == 200
            public_json = public_info.json()
            assert public_json["start_at"].startswith(window_payload["start_at"])
            assert public_json["end_at"].startswith(window_payload["end_at"])
            assert "title" not in public_json and "open_until" not in public_json
            assert not {"id", "participants", "entries", "entry_count", "public_token", "public_path"}.intersection(public_json)
            external_payload = {
                "name": "临时参与者",
                "student_id": "26000203",
                "college": "computer-science-technology",
            }
            external_doc = client.post(f"/api/daily-leave/public/{token}/document", json=external_payload)
            assert external_doc.status_code == 200, external_doc.text
            external_repeat = client.post(
                f"/api/daily-leave/public/{token}/document",
                json={**external_payload, "name": "其他输入"},
            )
            assert external_repeat.status_code == 200
            _, repeat_table = docx_text(external_repeat.content)
            assert repeat_table[0][1] == ["临时参与者", "26000203", "计算机科学与技术学院"]
            public_entry = db.scalar(
                select(DailyLeaveEntry).where(
                    DailyLeaveEntry.window_id == window_id,
                    DailyLeaveEntry.student_id_snapshot == external_payload["student_id"],
                )
            )
            assert public_entry is not None and public_entry.participant_type == "other"
            public_paragraphs, public_tables = docx_text(external_doc.content)
            assert public_tables[0][1] == ["临时参与者", "26000203", "计算机科学与技术学院"]
            assert "entries" not in "".join(public_paragraphs)
            assert client.post(
                f"/api/daily-leave/public/{token}/document",
                json={**external_payload, "student_id": "１２３４５６７８"},
            ).status_code == 422
            assert client.post(
                f"/api/daily-leave/public/{token}/document",
                json={**external_payload, "college": "invalid-college"},
            ).status_code == 422
            assert client.post(
                f"/api/daily-leave/public/{token}/document",
                json={**external_payload, "start_at": window_payload["start_at"]},
            ).status_code == 422

            # Re-downloading keeps the first member identity snapshot despite later Profile edits.
            formal.name = "后来改名"
            formal.student_id = "26000991"
            formal.college = "art"
            formal.team_membership = "reserve"
            db.commit()
            set_actor(formal)
            repeated_doc = client.post(f"/api/daily-leave/windows/{window_id}/document")
            assert repeated_doc.status_code == 200
            repeated_paragraphs, repeated_tables = docx_text(repeated_doc.content)
            assert repeated_tables[0][1] == signed_tables[0][1]
            assert db.scalar(select(func.count(DailyLeaveEntry.id)).where(DailyLeaveEntry.window_id == window_id)) == 4

            # Closing or reaching the deadline blocks both signed-in and public generation.
            set_actor(admin)
            closed = client.post(f"/api/daily-leave/admin/windows/{window_id}/close")
            assert closed.status_code == 200 and closed.json()["status"] == "closed"
            set_actor(formal)
            assert client.post(f"/api/daily-leave/windows/{window_id}/document").status_code == 409
            set_actor(None)
            assert client.post(f"/api/daily-leave/public/{token}/document", json=external_payload).status_code == 409
            assert client.get(f"/api/daily-leave/public/{token}").status_code == 404
            assert client.get("/api/daily-leave/public/not-a-valid-token").status_code == 404

            set_actor(admin)
            expired_payload = {
                "title": "已截止的窗口",
                "start_at": (now + timedelta(days=5)).replace(hour=13, minute=30).isoformat(timespec="minutes"),
                "end_at": (now + timedelta(days=5)).replace(hour=17, minute=10).isoformat(timespec="minutes"),
            }
            expired_response = client.post("/api/daily-leave/admin/windows", json=expired_payload)
            assert expired_response.status_code == 201, expired_response.text
            expired_window = expired_response.json()
            window_ids.add(expired_window["id"])
            expired_row = db.get(DailyLeaveWindow, expired_window["id"])
            assert expired_row is not None
            expired_row.public_enabled = True
            expired_row.public_token = "expired-test-token-123456789012345678901234"
            expired_row.start_at = local_now() - timedelta(hours=2)
            expired_row.end_at = local_now() - timedelta(hours=1)
            expired_row.open_until = expired_row.end_at
            db.commit()
            expired_token = expired_row.public_token
            set_actor(None)
            assert client.get(f"/api/daily-leave/public/{expired_token}").status_code == 404
            assert client.post(
                f"/api/daily-leave/public/{expired_token}/document",
                json={"name": "截止后报名", "student_id": "26000204", "college": "art"},
            ).status_code == 409

            # Camp docs use immutable college snapshots, with one separate list for each college.
            set_actor(admin)
            camp_payload = {
                "title": "测试冬令营活动",
                "type": "winter",
                "start_date": (date.today() + timedelta(days=20)).isoformat(),
                "end_date": (date.today() + timedelta(days=25)).isoformat(),
                "collection_deadline": (local_now() + timedelta(days=3)).isoformat(timespec="minutes"),
            }
            camp_response = client.post("/api/camp-leave/admin/events", json=camp_payload)
            assert camp_response.status_code == 201, camp_response.text
            camp_id = camp_response.json()["id"]
            event_ids.add(camp_id)
            set_actor(reserve)
            assert client.post(f"/api/camp-leave/events/{camp_id}/join").status_code == 201
            set_actor(formal)
            # Use the original formal snapshot from before Profile was edited by restoring only for camp signup.
            formal.name = "正式成员甲"
            formal.student_id = "26000201"
            formal.college = "electronic-science-engineering"
            formal.team_membership = "formal"
            db.commit()
            assert client.post(f"/api/camp-leave/events/{camp_id}/join").status_code == 201
            set_actor(admin)
            assert client.post(f"/api/camp-leave/admin/events/{camp_id}/close").status_code == 200
            signed_camp = client.get(
                f"/api/camp-leave/admin/events/{camp_id}/colleges/electronic-science-engineering/document"
            )
            assert signed_camp.status_code == 200, signed_camp.text
            camp_paragraphs, camp_tables = docx_text(signed_camp.content)
            visible_camp_paragraphs = [line.replace("\u2060", "") for line in camp_paragraphs]
            assert any("吉甲大师双创基地机器人战队寒假创新实践活动" in line for line in visible_camp_paragraphs)
            assert not any("测试冬令营活动" in line for line in camp_paragraphs)
            assert any("日常管理" in line and "全天候的实验室活动安排" in line for line in camp_paragraphs)
            assert camp_tables[0][0] == ["序号", "姓名", "学号", "学院"]
            assert len(camp_tables[0]) == 2
            assert camp_tables[0][1][1:] == ["正式成员甲", "26000201", "电子科学与工程学院"]
            assert {hashlib.sha256(value).hexdigest() for value in package_media(signed_camp.content)} >= {
                hashlib.sha256(signature).hexdigest(), hashlib.sha256(seal).hexdigest()
            }
            other_college = client.get(
                f"/api/camp-leave/admin/events/{camp_id}/colleges/artificial-intelligence/document"
            )
            assert other_college.status_code == 200
            _, other_table = docx_text(other_college.content)
            assert len(other_table[0]) == 2
            assert other_table[0][1][1:] == ["Reserve Member", "26000202", "人工智能学院"]
            member_doc = client.get(f"/api/camp-leave/events/{camp_id}/document")
            assert member_doc.status_code == 404  # Admin has not joined as a camp participant.
            set_actor(reserve)
            member_doc = client.get(f"/api/camp-leave/events/{camp_id}/document")
            assert member_doc.status_code == 200
            member_paragraphs, member_tables = docx_text(member_doc.content)
            assert len(member_tables[0]) == 2
            assert member_tables[0][1][1:] == ["Reserve Member", "26000202", "人工智能学院"]
            assert any("寒假" in line for line in member_paragraphs)
            set_actor(admin)
            offline_camp = client.get(
                f"/api/camp-leave/admin/events/{camp_id}/colleges/electronic-science-engineering/document?offline=true"
            )
            assert offline_camp.status_code == 200
            assert package_media(offline_camp.content) == set()
            offline_camp_paragraphs, offline_camp_tables = docx_text(offline_camp.content)
            assert [paragraph.replace("\u00a0", "") for paragraph in offline_camp_paragraphs] == camp_paragraphs
            assert offline_camp_tables == camp_tables
            summer_content = build_camp_leave_college_docx(
                title="测试夏令营活动",
                event_type="summer",
                start_date=date(2027, 7, 1),
                end_date=date(2027, 7, 8),
                participants=[("夏令营学生", "26000299", "人工智能学院")],
                offline=True,
            )
            summer_paragraphs, summer_tables = docx_text(summer_content)
            visible_summer_paragraphs = [line.replace("\u2060", "") for line in summer_paragraphs]
            assert any("吉甲大师双创基地机器人战队暑假创新实践活动" in line for line in visible_summer_paragraphs)
            assert not any("测试夏令营活动" in line for line in summer_paragraphs)
            assert summer_tables[0][1][1:] == ["夏令营学生", "26000299", "人工智能学院"]
            assert package_media(summer_content) == set()

            # Existing aggregation tables, status model, and read-only history endpoint remain usable.
            legacy_run = SchoolLeaveRun(
                collected_at=local_now(),
                created_by=admin.id,
                reason="legacy history fixture",
                status="ready",
            )
            db.add(legacy_run)
            db.flush()
            legacy_run_ids.add(legacy_run.id)
            db.add(
                SchoolLeaveRequest(
                    member_id=admin.id,
                    start_at=now + timedelta(days=1),
                    end_at=now + timedelta(days=1, hours=2),
                    member_name_snapshot="历史申请人",
                    student_id_snapshot="26000290",
                    status="included",
                    run_id=legacy_run.id,
                )
            )
            db.commit()
            set_actor(admin)
            history_list = client.get("/api/school-leave/admin/runs")
            assert history_list.status_code == 200
            assert any(run["id"] == legacy_run.id for run in history_list.json())
            history_doc = client.get(f"/api/school-leave/admin/runs/{legacy_run.id}/history-document")
            assert history_doc.status_code == 200, history_doc.text
            assert Document(io.BytesIO(history_doc.content)).tables
        finally:
            app.dependency_overrides.clear()
            app.dependency_overrides.update(previous_overrides)
            if window_ids:
                db.execute(delete(DailyLeaveEntry).where(DailyLeaveEntry.window_id.in_(window_ids)))
                db.execute(delete(DailyLeaveWindow).where(DailyLeaveWindow.id.in_(window_ids)))
            db.execute(
                    delete(DailyLeaveEntry).where(
                    DailyLeaveEntry.member_id.in_([admin.id, formal.id, reserve.id, unassigned.id, incomplete.id])
                )
            )
            if event_ids:
                db.execute(delete(CampLeaveParticipant).where(CampLeaveParticipant.event_id.in_(event_ids)))
                db.execute(delete(CampLeaveEvent).where(CampLeaveEvent.id.in_(event_ids)))
            if legacy_run_ids:
                db.execute(delete(SchoolLeaveRequest).where(SchoolLeaveRequest.run_id.in_(legacy_run_ids)))
                db.execute(delete(SchoolLeaveRun).where(SchoolLeaveRun.id.in_(legacy_run_ids)))
            db.execute(delete(Member).where(Member.email.in_([admin_email, formal_email, reserve_email, unassigned_email, incomplete_email])))
            db.commit()
            if original_phone is None:
                os.environ.pop("LEAVE_CONTACT_PHONE", None)
            else:
                os.environ["LEAVE_CONTACT_PHONE"] = original_phone

    print("Leave V2 workflow and DOCX tests passed")


if __name__ == "__main__":
    main()
