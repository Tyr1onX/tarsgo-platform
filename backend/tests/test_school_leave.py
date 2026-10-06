"""MySQL-backed School Leave workflow, delivery, deletion, and privacy coverage."""

from __future__ import annotations

import io
import os
from datetime import datetime
from urllib.parse import unquote
from concurrent.futures import ThreadPoolExecutor

from docx import Document
from fastapi.testclient import TestClient
from sqlalchemy import delete, select

from app.auth import get_current_member
from app.db import SessionLocal, get_db
from app.main import app
from app.models import Member, SchoolLeaveRequest, SchoolLeaveRun
from app.school_leave import (
    SCHOOL_LEAVE_TEMPLATE_PATH,
    collect_pending_school_leave,
    format_school_leave_course_period,
    school_leave_run_document_filename,
)


TEST_EMAILS = (
    "leave-admin@example.com",
    "leave-member-d@example.com",
    "leave-a@example.com",
    "leave-b@example.com",
    "leave-c@example.com",
    "leave-missing@example.com",
)


def set_actor(member: Member) -> None:
    app.dependency_overrides[get_current_member] = lambda: member


def request_payload(start: str, end: str) -> dict[str, str]:
    return {"start_at": start, "end_at": end}


def open_docx(content: bytes) -> Document:
    return Document(io.BytesIO(content))


def docx_text(content: bytes) -> str:
    document = open_docx(content)
    text_parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            text_parts.extend(cell.text for cell in row.cells)
    return "\n".join(text_parts)


def docx_page_break_count(content: bytes) -> int:
    document = open_docx(content)
    return document._element.xml.count('w:type="page"')


def table_text(document: Document, index: int) -> str:
    return "\n".join(
        cell.text
        for row in document.tables[index].rows
        for cell in row.cells
    )


def main() -> None:
    assert format_school_leave_course_period(
        datetime(2026, 10, 8, 8, 0), datetime(2026, 10, 8, 11, 40)
    ) == "上午课程"
    assert format_school_leave_course_period(
        datetime(2026, 10, 8, 13, 0), datetime(2026, 10, 8, 17, 0)
    ) == "下午课程"
    assert format_school_leave_course_period(
        datetime(2026, 10, 8, 18, 0), datetime(2026, 10, 8, 19, 0)
    ) == "晚间课程"
    assert format_school_leave_course_period(
        datetime(2026, 10, 8, 10, 0), datetime(2026, 10, 8, 14, 0)
    ) == "当日对应课程"
    assert format_school_leave_course_period(
        datetime(2026, 10, 8, 20, 0), datetime(2026, 10, 9, 9, 0)
    ) == "当日对应课程"

    previous_phone = os.environ.get("LEAVE_CONTACT_PHONE")
    previous_cutoff = os.environ.get("LEAVE_DAILY_CUTOFF")
    previous_overrides = dict(app.dependency_overrides)

    with SessionLocal() as db:
        template_body = next(
            paragraph.text
            for paragraph in Document(SCHOOL_LEAVE_TEMPLATE_PATH).paragraphs
            if paragraph.text.startswith("以下学生因")
        )
        assert template_body == (
            "以下学生因参加2026 年 10 月 6 日 13:00 至 17:00的吉甲大师双创基地机器人战队创新实践活动，"
            "不能参加下午课程，特此证明。"
        )

        existing_ids = list(db.scalars(select(Member.id).where(Member.email.in_(TEST_EMAILS))))
        if existing_ids:
            run_ids = list(
                db.scalars(
                    select(SchoolLeaveRequest.run_id)
                    .where(
                        SchoolLeaveRequest.member_id.in_(existing_ids),
                        SchoolLeaveRequest.run_id.is_not(None),
                    )
                    .distinct()
                )
            )
            db.execute(delete(SchoolLeaveRequest).where(SchoolLeaveRequest.member_id.in_(existing_ids)))
            if run_ids:
                db.execute(delete(SchoolLeaveRun).where(SchoolLeaveRun.id.in_(run_ids)))
            db.execute(delete(SchoolLeaveRun).where(SchoolLeaveRun.created_by.in_(existing_ids)))
            db.execute(delete(Member).where(Member.id.in_(existing_ids)))
            db.commit()

        admin = Member(
            name="测试管理员",
            email=TEST_EMAILS[0],
            student_id="TEST900001",
            role="admin",
            status="active",
        )
        member_d = Member(
            name="测试成员丁",
            email=TEST_EMAILS[1],
            student_id="TEST900002",
            role="member",
            status="active",
        )
        member_a = Member(
            name="测试甲",
            email=TEST_EMAILS[2],
            student_id="TEST100001",
            role="member",
            status="active",
        )
        member_b = Member(
            name="测试乙",
            email=TEST_EMAILS[3],
            student_id="TEST100002",
            role="member",
            status="active",
        )
        member_c = Member(
            name="测试丙",
            email=TEST_EMAILS[4],
            student_id="TEST100003",
            role="member",
            status="active",
        )
        missing_id = Member(
            name="测试未填学号",
            email=TEST_EMAILS[5],
            student_id=None,
            role="member",
            status="active",
        )
        db.add_all([admin, member_d, member_a, member_b, member_c, missing_id])
        db.commit()
        for member in (admin, member_d, member_a, member_b, member_c, missing_id):
            db.refresh(member)

        def override_database():
            yield db

        app.dependency_overrides[get_db] = override_database
        client = TestClient(app)

        try:
            # A student id is mandatory for school-leave submission.
            set_actor(missing_id)
            response = client.post(
                "/api/school-leave/requests",
                json=request_payload("2026-10-08T13:00:00", "2026-10-08T17:00:00"),
            )
            assert response.status_code == 409, response.text
            assert response.json()["detail"] == "请先完善学号，生成学校请假材料时需要使用。"

            # Task member summaries must not start exposing student ids.
            set_actor(admin)
            task_assignees = client.get("/api/tasks/assignees")
            assert task_assignees.status_code == 200, task_assignees.text
            assert all("student_id" not in row for row in task_assignees.json())

            # A member can update/withdraw only their own pending request.
            set_actor(member_a)
            first = client.post(
                "/api/school-leave/requests",
                json=request_payload("2026-10-08T13:00:00", "2026-10-08T17:00:00"),
            )
            assert first.status_code == 201, first.text
            first_id = first.json()["id"]
            assert first.json()["student_id_snapshot"] == "TEST100001"
            assert first.json()["run_status"] is None

            updated = client.patch(
                f"/api/school-leave/requests/{first_id}",
                json=request_payload("2026-10-08T13:05:00", "2026-10-08T17:00:00"),
            )
            assert updated.status_code == 200, updated.text
            assert updated.json()["start_at"].startswith("2026-10-08T13:05")

            set_actor(member_b)
            forbidden_update = client.patch(
                f"/api/school-leave/requests/{first_id}",
                json=request_payload("2026-10-08T13:10:00", "2026-10-08T17:00:00"),
            )
            assert forbidden_update.status_code == 404

            set_actor(member_a)
            withdrawn = client.post(f"/api/school-leave/requests/{first_id}/withdraw")
            assert withdrawn.status_code == 200
            assert withdrawn.json()["status"] == "withdrawn"
            assert withdrawn.json()["run_status"] is None

            # Prepare one exact shared group plus one different group.
            a_response = client.post(
                "/api/school-leave/requests",
                json=request_payload("2026-10-08T13:00:00", "2026-10-08T17:00:00"),
            )
            assert a_response.status_code == 201
            a_request_id = a_response.json()["id"]

            set_actor(member_b)
            b_response = client.post(
                "/api/school-leave/requests",
                json=request_payload("2026-10-08T13:00:00", "2026-10-08T17:00:00"),
            )
            assert b_response.status_code == 201
            b_request_id = b_response.json()["id"]

            set_actor(member_c)
            c_response = client.post(
                "/api/school-leave/requests",
                json=request_payload("2026-10-08T15:00:00", "2026-10-08T17:00:00"),
            )
            assert c_response.status_code == 201
            c_request_id = c_response.json()["id"]

            # Members never gain school-leave administration permission.
            set_actor(member_d)
            assert client.get("/api/school-leave/admin/runs").status_code == 403
            assert client.post("/api/school-leave/admin/runs/collect").status_code == 403
            set_actor(member_a)
            assert client.get("/api/school-leave/admin/requests").status_code == 403

            # Collection groups only exact equal start/end times.
            set_actor(admin)
            collected = client.post("/api/school-leave/admin/runs/collect")
            assert collected.status_code == 200, collected.text
            run_one = collected.json()
            assert run_one is not None
            run_one_id = run_one["id"]
            assert "send_message" not in run_one
            assert [(group["time_text"], group["count"]) for group in run_one["groups"]] == [
                ("2026 年 10 月 8 日 13:00 至 17:00", 2),
                ("2026 年 10 月 8 日 15:00 至 17:00", 1),
            ]

            # Members can read only the run status attached to their own requests.
            set_actor(member_a)
            ready_requests = client.get("/api/school-leave/requests")
            assert ready_requests.status_code == 200, ready_requests.text
            ready_request = next(item for item in ready_requests.json() if item["id"] == a_request_id)
            assert ready_request["status"] == "included"
            assert ready_request["run_status"] == "ready"
            assert not {"reason", "groups", "created_by", "downloaded_by"} & ready_request.keys()
            assert client.get("/api/school-leave/admin/runs").status_code == 403

            # New document/delete endpoints remain admin-only.
            set_actor(member_d)
            assert client.get(f"/api/school-leave/admin/runs/{run_one_id}/document").status_code == 403
            assert client.delete(f"/api/school-leave/admin/runs/{run_one_id}").status_code == 403
            set_actor(member_a)
            assert client.get(f"/api/school-leave/admin/runs/{run_one_id}/document").status_code == 403
            assert client.delete(f"/api/school-leave/admin/runs/{run_one_id}").status_code == 403

            # Included requests are frozen for members.
            included_update = client.patch(
                f"/api/school-leave/requests/{a_request_id}",
                json=request_payload("2026-10-08T13:05:00", "2026-10-08T17:00:00"),
            )
            assert included_update.status_code == 409

            # Profile changes never mutate frozen snapshots.
            member_a.student_id = "TEST199999"
            db.commit()
            set_actor(admin)
            run_one_fresh = next(
                run for run in client.get("/api/school-leave/admin/runs").json()
                if run["id"] == run_one_id
            )
            shared_group_ids = [person["student_id"] for person in run_one_fresh["groups"][0]["members"]]
            assert "TEST100001" in shared_group_ids
            assert "TEST199999" not in shared_group_ids

            # A ready run cannot be deleted; it must go through cancellation.
            ready_delete = client.delete(f"/api/school-leave/admin/runs/{run_one_id}")
            assert ready_delete.status_code == 409
            assert "取消" in ready_delete.json()["detail"]

            # Later submissions stay pending; cancelling a ready run releases included rows.
            set_actor(member_a)
            later = client.post(
                "/api/school-leave/requests",
                json=request_payload("2026-10-08T18:00:00", "2026-10-08T19:00:00"),
            )
            assert later.status_code == 201
            later_id = later.json()["id"]
            assert later.json()["status"] == "pending"
            assert later.json()["run_id"] is None

            set_actor(admin)
            cancelled = client.post(f"/api/school-leave/admin/runs/{run_one_id}/cancel")
            assert cancelled.status_code == 200, cancelled.text
            assert cancelled.json()["status"] == "cancelled"
            assert client.get(f"/api/school-leave/admin/runs/{run_one_id}/document").status_code == 409
            db.expire_all()
            for request_id in (a_request_id, b_request_id, c_request_id):
                request = db.get(SchoolLeaveRequest, request_id)
                assert request and request.status == "pending" and request.run_id is None

            # Cancelled history can be cleaned without touching released requests.
            cancelled_delete = client.delete(f"/api/school-leave/admin/runs/{run_one_id}")
            assert cancelled_delete.status_code == 204, cancelled_delete.text
            db.expire_all()
            assert db.get(SchoolLeaveRun, run_one_id) is None
            for request_id in (a_request_id, b_request_id, c_request_id):
                assert db.get(SchoolLeaveRequest, request_id) is not None

            # Re-collect released requests plus the later supplement.
            set_actor(member_d)
            morning = client.post(
                "/api/school-leave/requests",
                json=request_payload("2026-10-08T08:00:00", "2026-10-08T09:00:00"),
            )
            assert morning.status_code == 201, morning.text
            morning_request_id = morning.json()["id"]

            missing_id.student_id = "99000006"
            db.commit()
            set_actor(missing_id)
            cross_period = client.post(
                "/api/school-leave/requests",
                json=request_payload("2026-10-08T10:00:00", "2026-10-08T14:00:00"),
            )
            assert cross_period.status_code == 201, cross_period.text
            cross_period_request_id = cross_period.json()["id"]

            # Re-collect released requests plus the later supplement and period examples.
            recollected = client.post("/api/school-leave/admin/runs/collect")
            assert recollected.status_code == 200, recollected.text
            run_two = recollected.json()
            run_two_id = run_two["id"]
            assert run_two["request_count"] == 6
            assert [group["count"] for group in run_two["groups"]] == [1, 1, 2, 1, 1]
            db.expire_all()
            assert db.get(SchoolLeaveRequest, morning_request_id).run_id == run_two_id
            assert db.get(SchoolLeaveRequest, cross_period_request_id).run_id == run_two_id
            assert "send_message" not in run_two

            # Stored batch reason remains editable, while the form wording stays fixed.
            reason = "参加虚构机器人战队校内创新实践活动"
            changed_reason = client.patch(
                f"/api/school-leave/admin/runs/{run_two_id}/reason",
                json={"reason": reason},
            )
            assert changed_reason.status_code == 200
            assert changed_reason.json()["reason"] == reason

            # Missing private phone configuration safely blocks the run document.
            os.environ.pop("LEAVE_CONTACT_PHONE", None)
            missing_phone = client.get(f"/api/school-leave/admin/runs/{run_two_id}/document")
            assert missing_phone.status_code == 503
            assert "LEAVE_CONTACT_PHONE" in missing_phone.json()["detail"]

            # Multi-group run: one DOCX, one full leave form per group, page-break separated.
            os.environ["LEAVE_CONTACT_PHONE"] = "000-0000-0000"
            doc_response = client.get(f"/api/school-leave/admin/runs/{run_two_id}/document")
            assert doc_response.status_code == 200, doc_response.text
            assert doc_response.headers["content-type"].startswith(
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            )
            content_disposition = doc_response.headers["content-disposition"]
            assert ".docx" in content_disposition
            run_two_row = db.get(SchoolLeaveRun, run_two_id)
            assert run_two_row is not None
            assert unquote(content_disposition.split("UTF-8''", 1)[1]) == school_leave_run_document_filename(
                db,
                run_two_row,
            )
            rendered = docx_text(doc_response.content)
            assert rendered.count("请假条") == 5
            assert reason not in rendered
            assert rendered.count("吉甲大师双创基地机器人战队创新实践活动") == 5
            assert "以下学生因参加2026 年 10 月 8 日 08:00 至 09:00的吉甲大师双创基地机器人战队创新实践活动，不能参加上午课程，特此证明。" in rendered
            assert "以下学生因参加2026 年 10 月 8 日 10:00 至 14:00的吉甲大师双创基地机器人战队创新实践活动，不能参加当日对应课程，特此证明。" in rendered
            assert "以下学生因参加2026 年 10 月 8 日 13:00 至 17:00的吉甲大师双创基地机器人战队创新实践活动，不能参加下午课程，特此证明。" in rendered
            assert "以下学生因参加2026 年 10 月 8 日 18:00 至 19:00的吉甲大师双创基地机器人战队创新实践活动，不能参加晚间课程，特此证明。" in rendered
            assert "2026 年 10 月 8 日 13:00 至 17:00" in rendered
            assert "2026 年 10 月 8 日 15:00 至 17:00" in rendered
            assert "2026 年 10 月 8 日 18:00 至 19:00" in rendered
            assert "联系电话：000-0000-0000" in rendered
            document = open_docx(doc_response.content)
            template_document = Document(SCHOOL_LEAVE_TEMPLATE_PATH)
            assert len(document.tables) == 5
            assert docx_page_break_count(doc_response.content) == 4
            assert document.sections[0].top_margin == template_document.sections[0].top_margin
            assert document.sections[0].bottom_margin == template_document.sections[0].bottom_margin
            assert document.sections[0].left_margin == template_document.sections[0].left_margin
            assert document.sections[0].right_margin == template_document.sections[0].right_margin
            assert document.tables[0].autofit == template_document.tables[0].autofit
            assert [
                cell._tc.tcPr.tcW.w for cell in document.tables[0].rows[0].cells
            ] == [
                cell._tc.tcPr.tcW.w for cell in template_document.tables[0].rows[0].cells
            ]
            assert [cell.text for cell in document.tables[0].rows[0].cells] == ["姓名", "学号"]
            assert len(document.tables[0].rows) == 3
            morning_group = table_text(document, 0)
            cross_period_group = table_text(document, 1)
            first_group = table_text(document, 2)
            second_group = table_text(document, 3)
            third_group = table_text(document, 4)
            assert "测试成员丁" in morning_group and "TEST900002" in morning_group
            assert "测试未填学号" in cross_period_group and "99000006" in cross_period_group
            assert "测试甲" in first_group and "TEST100001" in first_group
            assert "测试乙" in first_group and "TEST100002" in first_group
            assert "测试丙" not in first_group
            assert "测试丙" in second_group and "TEST100003" in second_group
            assert "测试甲" not in second_group and "测试乙" not in second_group
            assert "测试甲" in third_group and "TEST199999" in third_group
            assert "TEST100001" not in third_group
            assert "测试乙" not in third_group and "测试丙" not in third_group

            # Backward-compatible single-group endpoint remains available, but ZIP is gone.
            group_doc = client.get(f"/api/school-leave/admin/runs/{run_two_id}/documents/0")
            assert group_doc.status_code == 200
            assert len(open_docx(group_doc.content).tables) == 1
            assert client.get(f"/api/school-leave/admin/runs/{run_two_id}/documents.zip").status_code == 404

            # A successful first download starts processing automatically; no manual sent action remains.
            db.expire_all()
            downloaded_run = db.get(SchoolLeaveRun, run_two_id)
            assert downloaded_run is not None
            assert downloaded_run.status == "awaiting_return"
            assert downloaded_run.downloaded_by == admin.id
            assert downloaded_run.downloaded_at is not None
            first_downloaded_at = downloaded_run.downloaded_at
            assert client.post(f"/api/school-leave/admin/runs/{run_two_id}/sent").status_code == 404
            assert client.get(f"/api/school-leave/admin/runs/{run_two_id}/document").status_code == 200
            db.expire_all()
            downloaded_run = db.get(SchoolLeaveRun, run_two_id)
            assert downloaded_run is not None
            assert downloaded_run.status == "awaiting_return"
            assert downloaded_run.downloaded_at == first_downloaded_at
            assert client.post(f"/api/school-leave/admin/runs/{run_two_id}/cancel").status_code == 409
            assert client.patch(
                f"/api/school-leave/admin/runs/{run_two_id}/reason",
                json={"reason": "不应允许修改"},
            ).status_code == 409

            # Members see processing status on their own requests but gain no management access.
            set_actor(member_a)
            processing_requests = client.get("/api/school-leave/requests")
            assert processing_requests.status_code == 200, processing_requests.text
            processing_request = next(item for item in processing_requests.json() if item["id"] == a_request_id)
            assert processing_request["status"] == "included"
            assert processing_request["run_status"] == "awaiting_return"
            assert processing_request["result_state"] is None
            assert client.get("/api/school-leave/admin/runs").status_code == 403

            # Deleting processing history remains admin-only.
            set_actor(member_d)
            assert client.delete(f"/api/school-leave/admin/runs/{run_two_id}").status_code == 403
            set_actor(member_a)
            assert client.delete(f"/api/school-leave/admin/runs/{run_two_id}").status_code == 403

            # A later request forms a fresh supplementary run and never mutates the sent one.
            set_actor(member_b)
            supplement = client.post(
                "/api/school-leave/requests",
                json=request_payload("2026-10-08T20:00:00", "2026-10-08T21:00:00"),
            )
            assert supplement.status_code == 201
            supplement_id = supplement.json()["id"]

            set_actor(admin)
            supplement_run = client.post("/api/school-leave/admin/runs/collect")
            assert supplement_run.status_code == 200
            run_three = supplement_run.json()
            run_three_id = run_three["id"]
            assert run_three_id != run_two_id
            assert run_three["request_count"] == 1
            db.expire_all()
            sent_request_ids = {
                request.id
                for request in db.scalars(
                    select(SchoolLeaveRequest).where(SchoolLeaveRequest.run_id == run_two_id)
                )
            }
            assert supplement_id not in sent_request_ids
            assert db.get(SchoolLeaveRun, run_two_id).status == "awaiting_return"

            # One-group run produces one DOCX with no page break.
            single_doc = client.get(f"/api/school-leave/admin/runs/{run_three_id}/document")
            assert single_doc.status_code == 200
            single_document = open_docx(single_doc.content)
            assert len(single_document.tables) == 1
            assert len(single_document.paragraphs) == len(template_document.paragraphs)
            assert docx_page_break_count(single_doc.content) == 0
            assert "2026 年 10 月 8 日 20:00 至 21:00" in docx_text(single_doc.content)
            assert "测试乙" in table_text(single_document, 0)
            assert client.delete(f"/api/school-leave/admin/runs/{run_three_id}").status_code == 409

            # Completed history deletion removes the run and every linked request, never orphaning included rows.
            run_two_row = db.get(SchoolLeaveRun, run_two_id)
            assert run_two_row is not None
            run_two_row.status = "completed"
            db.commit()
            run_two_request_ids = list(
                db.scalars(select(SchoolLeaveRequest.id).where(SchoolLeaveRequest.run_id == run_two_id))
            )
            assert len(run_two_request_ids) == 6
            sent_delete = client.delete(f"/api/school-leave/admin/runs/{run_two_id}")
            assert sent_delete.status_code == 204, sent_delete.text
            db.expire_all()
            assert db.get(SchoolLeaveRun, run_two_id) is None
            assert all(db.get(SchoolLeaveRequest, request_id) is None for request_id in run_two_request_ids)
            orphaned = list(
                db.scalars(
                    select(SchoolLeaveRequest).where(
                        SchoolLeaveRequest.status == "included",
                        SchoolLeaveRequest.run_id.is_(None),
                    )
                )
            )
            assert orphaned == []
            assert client.delete(f"/api/school-leave/admin/runs/{run_two_id}").status_code == 404

            # Deletion is atomic: a commit failure rolls back request and run deletion together.
            run_three_row = db.get(SchoolLeaveRun, run_three_id)
            assert run_three_row is not None
            run_three_row.status = "completed"
            db.commit()
            original_commit = db.commit

            def failing_commit() -> None:
                raise RuntimeError("forced delete commit failure")

            db.commit = failing_commit  # type: ignore[method-assign]
            failing_client = TestClient(app, raise_server_exceptions=False)
            failed_delete = failing_client.delete(f"/api/school-leave/admin/runs/{run_three_id}")
            assert failed_delete.status_code == 500
            db.commit = original_commit  # type: ignore[method-assign]
            db.rollback()
            db.expire_all()
            assert db.get(SchoolLeaveRun, run_three_id) is not None
            supplement_row = db.get(SchoolLeaveRequest, supplement_id)
            assert supplement_row is not None
            assert supplement_row.status == "included"
            assert supplement_row.run_id == run_three_id

            # After restoring the transaction, the same completed history can be deleted normally.
            deleted_three = client.delete(f"/api/school-leave/admin/runs/{run_three_id}")
            assert deleted_three.status_code == 204
            db.expire_all()
            assert db.get(SchoolLeaveRun, run_three_id) is None
            assert db.get(SchoolLeaveRequest, supplement_id) is None

            # No pending requests is an idempotent safe no-op.
            assert collect_pending_school_leave(db, created_by=None) is None

            # Two simultaneous collectors can include a pending request only once.
            set_actor(member_c)
            concurrency_request = client.post(
                "/api/school-leave/requests",
                json=request_payload("2026-10-09T08:00:00", "2026-10-09T09:00:00"),
            )
            assert concurrency_request.status_code == 201
            concurrency_request_id = concurrency_request.json()["id"]
            db.commit()

            def collect_worker() -> int | None:
                with SessionLocal() as worker_db:
                    run = collect_pending_school_leave(worker_db, created_by=None)
                    return run.id if run else None

            with ThreadPoolExecutor(max_workers=2) as executor:
                run_ids = list(executor.map(lambda _: collect_worker(), range(2)))

            created_run_ids = [run_id for run_id in run_ids if run_id is not None]
            assert len(created_run_ids) == 1, run_ids
            db.expire_all()
            concurrency_row = db.get(SchoolLeaveRequest, concurrency_request_id)
            assert concurrency_row and concurrency_row.status == "included"
            assert concurrency_row.run_id == created_run_ids[0]

            # Download filenames hide database run IDs and number only valid same-day supplements.
            filename_runs = [
                SchoolLeaveRun(
                    collected_at=datetime(2030, 1, 2, 9, 0),
                    created_by=admin.id,
                    reason="文件名测试",
                    status="completed",
                ),
                SchoolLeaveRun(
                    collected_at=datetime(2030, 1, 2, 10, 0),
                    created_by=admin.id,
                    reason="文件名测试",
                    status="cancelled",
                ),
                SchoolLeaveRun(
                    collected_at=datetime(2030, 1, 2, 11, 0),
                    created_by=admin.id,
                    reason="文件名测试",
                    status="ready",
                ),
                SchoolLeaveRun(
                    collected_at=datetime(2030, 1, 2, 12, 0),
                    created_by=admin.id,
                    reason="文件名测试",
                    status="completed",
                ),
                SchoolLeaveRun(
                    collected_at=datetime(2030, 1, 3, 9, 0),
                    created_by=admin.id,
                    reason="文件名测试",
                    status="ready",
                ),
            ]
            db.add_all(filename_runs)
            db.commit()
            assert school_leave_run_document_filename(db, filename_runs[0]) == "吉甲大师请假条_2030-01-02.docx"
            assert school_leave_run_document_filename(db, filename_runs[2]) == "吉甲大师请假条_2030-01-02_补充1.docx"
            assert school_leave_run_document_filename(db, filename_runs[3]) == "吉甲大师请假条_2030-01-02_补充2.docx"
            assert school_leave_run_document_filename(db, filename_runs[4]) == "吉甲大师请假条_2030-01-03.docx"

            # Cutoff parsing is configuration-only and visible to admins without exposing phone.
            os.environ["LEAVE_DAILY_CUTOFF"] = "11:30"
            set_actor(admin)
            config = client.get("/api/school-leave/admin/config")
            assert config.status_code == 200
            assert config.json() == {
                "daily_cutoff": "11:30",
                "contact_phone_configured": True,
            }

            # Sanity-check request statuses that intentionally remain after workflow cleanup.
            db.expire_all()
            assert db.get(SchoolLeaveRequest, first_id).status == "withdrawn"
            assert db.get(SchoolLeaveRequest, later_id) is None

        finally:
            app.dependency_overrides.clear()
            app.dependency_overrides.update(previous_overrides)
            test_member_ids = [
                admin.id,
                member_d.id,
                member_a.id,
                member_b.id,
                member_c.id,
                missing_id.id,
            ]
            run_ids = list(
                db.scalars(
                    select(SchoolLeaveRequest.run_id)
                    .where(
                        SchoolLeaveRequest.member_id.in_(test_member_ids),
                        SchoolLeaveRequest.run_id.is_not(None),
                    )
                    .distinct()
                )
            )
            db.execute(delete(SchoolLeaveRequest).where(SchoolLeaveRequest.member_id.in_(test_member_ids)))
            if run_ids:
                db.execute(delete(SchoolLeaveRun).where(SchoolLeaveRun.id.in_(run_ids)))
            db.execute(delete(SchoolLeaveRun).where(SchoolLeaveRun.created_by.in_(test_member_ids)))
            db.execute(delete(Member).where(Member.id.in_(test_member_ids)))
            db.commit()

    if previous_phone is None:
        os.environ.pop("LEAVE_CONTACT_PHONE", None)
    else:
        os.environ["LEAVE_CONTACT_PHONE"] = previous_phone
    if previous_cutoff is None:
        os.environ.pop("LEAVE_DAILY_CUTOFF", None)
    else:
        os.environ["LEAVE_DAILY_CUTOFF"] = previous_cutoff

    print("School leave tests passed")


if __name__ == "__main__":
    main()
